"""生成人员安全智能运营中心工业产品销售培训 PPT。"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "人员安全智能运营中心-工业产品销售培训.pptx"

W, H = 13.333, 7.5
BG = "0A1118"
PANEL = "121D26"
PANEL_2 = "172631"
STEEL = "263945"
LINE = "334A57"
WHITE = "F2F6F7"
MUTED = "97AAB4"
YELLOW = "F3AA2B"
TEAL = "28C7B7"
CYAN = "4DA9D7"
RED = "E26363"
GREEN = "49C486"
FONT = "Microsoft YaHei"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_bg(slide, color: str = BG):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(color)


def add_box(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True, width=1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def add_text(slide, text, x, y, w, h, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT, margin=0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=16, color=WHITE, bullet=False, spacing=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(lines):
        if isinstance(item, tuple):
            text, item_color, bold = item
        else:
            text, item_color, bold = item, color, False
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = ("•  " if bullet else "") + text
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(item_color)
        p.space_after = Pt(spacing)
    return box


def add_title(slide, title: str, subtitle: str = "", index: int | None = None):
    add_box(slide, 0.45, 0.34, 0.08, 0.48, fill=YELLOW, line=YELLOW, radius=False, width=0)
    add_text(slide, title, 0.68, 0.27, 10.9, 0.5, size=25, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.7, 0.78, 11.7, 0.32, size=10.5, color=MUTED)
    if index is not None:
        add_text(slide, f"{index:02d}", 12.08, 0.3, 0.72, 0.35, size=12, color=YELLOW, bold=True, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.16), Inches(12.82), Inches(1.16))
    line.line.color.rgb = rgb(LINE)
    line.line.width = Pt(1)


def add_footer(slide, number: int):
    add_text(slide, "人员安全智能运营中心｜销售培训资料", 0.52, 7.15, 5.5, 0.18, size=8.5, color="6F838E")
    add_text(slide, f"{number:02d}", 12.15, 7.12, 0.65, 0.2, size=8.5, color="6F838E", align=PP_ALIGN.RIGHT)


def add_notes(slide, text: str):
    try:
        slide.notes_slide.notes_text_frame.text = text
    except Exception:
        pass


def add_badge(slide, text, x, y, w, color=TEAL):
    add_box(slide, x, y, w, 0.34, fill=PANEL_2, line=color, radius=True, width=1)
    add_text(slide, text, x + 0.04, y + 0.04, w - 0.08, 0.22, size=9.5, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_metric(slide, value, label, x, y, w, color=YELLOW):
    add_box(slide, x, y, w, 1.02, fill=PANEL, line=LINE)
    add_text(slide, value, x + 0.12, y + 0.13, w - 0.24, 0.38, size=24, color=color, bold=True)
    add_text(slide, label, x + 0.12, y + 0.57, w - 0.24, 0.25, size=10.5, color=MUTED)


def add_capability_card(slide, index, title, desc, x, y, w, h, color=TEAL):
    add_box(slide, x, y, w, h, fill=PANEL, line=LINE)
    add_box(slide, x + 0.16, y + 0.16, 0.38, 0.38, fill=color, line=color, radius=True, width=0)
    add_text(slide, str(index), x + 0.16, y + 0.21, 0.38, 0.2, size=10, color=BG, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.65, y + 0.16, w - 0.82, 0.3, size=14, bold=True)
    add_text(slide, desc, x + 0.18, y + 0.68, w - 0.36, h - 0.82, size=10.5, color=MUTED)


def new_slide(prs, number, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, title, subtitle, number)
    add_footer(slide, number)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "人员安全智能运营中心｜工业产品销售培训"
    prs.core_properties.subject = "面向工业企业的人员安全、设备运维与智能运营平台"
    prs.core_properties.author = "人员安全智能运营中心"

    # 01 封面
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_box(slide, 0, 0, 4.5, H, fill="0D1921", line="0D1921", radius=False, width=0)
    for i in range(9):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.7 + i * 0.85), Inches(0), Inches(2.4 + i * 0.85), Inches(7.5))
        line.line.color.rgb = rgb("162A35")
        line.line.width = Pt(0.7)
    add_box(slide, 0.7, 0.7, 0.58, 0.58, fill=YELLOW, line=YELLOW, radius=True, width=0)
    add_text(slide, "AI", 0.7, 0.87, 0.58, 0.2, size=12, color=BG, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "人员安全智能运营中心", 0.72, 2.05, 8.8, 0.78, size=34, bold=True)
    add_text(slide, "面向工业现场的人、地、票、警、设备与生产一体化智能运营平台", 0.76, 2.95, 9.5, 0.46, size=17, color="C1D0D6")
    add_box(slide, 0.76, 3.65, 1.1, 0.05, fill=YELLOW, line=YELLOW, radius=False, width=0)
    add_text(slide, "让人员安全看得见，让设备风险早发现", 0.76, 3.95, 8.8, 0.42, size=20, color=YELLOW, bold=True)
    add_badge(slide, "工业产品销售培训", 0.76, 5.35, 1.9, YELLOW)
    add_badge(slide, "解决方案版", 2.85, 5.35, 1.35, TEAL)
    add_text(slide, "INTEGRATED INDUSTRIAL INTELLIGENCE", 0.78, 6.75, 5.5, 0.25, size=9, color="607985")
    add_notes(slide, "开场不要先讲技术。用一句话定位：我们把工业现场分散的人、作业和设备数据，变成管理者可以直接提问、直接决策的统一智能入口。")

    # 02 一句话定位
    slide = new_slide(prs, 2, "销售首先讲清楚：我们卖的是什么", "不要把产品介绍成聊天机器人，也不要只强调人员定位")
    add_text(slide, "一个入口", 0.72, 1.65, 3.2, 0.52, size=30, color=YELLOW, bold=True)
    add_text(slide, "连接工业现场全部关键数据", 0.75, 2.25, 3.4, 0.42, size=15, color=WHITE, bold=True)
    add_rich_lines(slide, ["自然语言查询", "网页与移动端", "统一身份与权限"], 0.75, 2.9, 3.25, 1.8, size=13, color=MUTED, bullet=True)
    arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(4.25), Inches(2.42), Inches(1.05), Inches(1.2))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = rgb(YELLOW); arrow.line.color.rgb = rgb(YELLOW)
    add_text(slide, "一套能力", 5.55, 1.65, 3.2, 0.52, size=30, color=TEAL, bold=True)
    add_text(slide, "查询、分析、预测、可视化", 5.58, 2.25, 3.45, 0.42, size=15, bold=True)
    add_rich_lines(slide, ["动态调用业务接口", "数据驱动结论", "本地与云端协同"], 5.58, 2.9, 3.2, 1.8, size=13, color=MUTED, bullet=True)
    add_box(slide, 9.25, 1.55, 3.2, 4.25, fill=PANEL_2, line=YELLOW)
    add_text(slide, "销售标准表述", 9.6, 1.9, 2.5, 0.4, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "“面向工业企业的人员安全与设备智能运营平台，以统一智能入口连接现场业务系统，帮助客户更快发现风险、更快获取结论、更快形成闭环。”", 9.67, 2.65, 2.35, 2.05, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_notes(slide, "强调三个关键词：工业现场、统一入口、风险闭环。客户问是不是大模型套壳时，回答：模型只是交互入口，真正价值在业务接口、数据关联、工具路由和安全闭环。")

    # 03 客户画像
    slide = new_slide(prs, 3, "目标客户与关键决策人", "先判断谁有痛点、谁有预算、谁影响选型")
    customers = [
        ("能源集团", "电力、热力、燃气、综合能源", YELLOW),
        ("流程工业", "化工、煤化工、冶金、建材", TEAL),
        ("大型园区", "工业园、制造基地、仓储物流", CYAN),
    ]
    for i, (name, desc, color) in enumerate(customers):
        x = 0.65 + i * 4.18
        add_capability_card(slide, i + 1, name, desc, x, 1.55, 3.75, 1.45, color)
    buyers = [
        ("主要购买者", "安全管理部｜生产运营部｜设备管理部｜数字化部门"),
        ("关键影响者", "厂长/总工程师｜信息中心｜承包商管理部门"),
        ("核心使用者", "调度员｜值班长｜安全员｜设备工程师｜现场负责人"),
    ]
    for i, (title, desc) in enumerate(buyers):
        y = 3.4 + i * 0.88
        add_box(slide, 0.68, y, 12.0, 0.68, fill=PANEL, line=LINE)
        add_text(slide, title, 0.95, y + 0.18, 1.55, 0.25, size=13, color=YELLOW if i == 0 else TEAL, bold=True)
        add_text(slide, desc, 2.75, y + 0.18, 9.45, 0.25, size=13, color=WHITE)
    add_notes(slide, "面向不同角色切换价值点：安全部门讲人员和作业风险，设备部门讲预测维护，信息化部门讲接口复用、本地部署和权限审计。")

    # 04 客户痛点
    slide = new_slide(prs, 4, "客户为什么需要这套产品", "不要从功能出发，要从现场管理成本和风险出发")
    pains = [
        ("系统多", "定位、门禁、作业票、设备监测彼此孤立", "跨系统查询慢"),
        ("数据散", "同一人员、区域和设备缺少统一关联", "信息难以互证"),
        ("报警多", "大量告警缺少优先级和连续性分析", "关键风险被淹没"),
        ("报表慢", "临时汇报依赖人工导出、拼表和解释", "管理响应滞后"),
        ("维护被动", "监测数据只看当前值，缺少趋势和寿命判断", "故障后处置"),
        ("数据敏感", "人员、生产和设备数据不适合全部上云", "智能化推进受阻"),
    ]
    for i, (title, desc, result) in enumerate(pains):
        col, row = i % 3, i // 3
        x, y = 0.64 + col * 4.18, 1.55 + row * 2.45
        add_box(slide, x, y, 3.78, 2.02, fill=PANEL, line=LINE)
        add_text(slide, f"0{i+1}", x + 0.2, y + 0.2, 0.5, 0.3, size=12, color=YELLOW, bold=True)
        add_text(slide, title, x + 0.8, y + 0.18, 2.6, 0.34, size=17, bold=True)
        add_text(slide, desc, x + 0.22, y + 0.72, 3.3, 0.52, size=11.5, color=MUTED)
        add_badge(slide, result, x + 0.22, y + 1.45, 1.55, RED)
    add_notes(slide, "请销售在客户现场至少确认三个痛点，不要把所有痛点一次讲完。客户对哪个问题反应最强，就从对应场景进入产品能力。")

    # 05 能力全景
    slide = new_slide(prs, 5, "六位一体的工业智能运营能力", "围绕人、地、票、警、设备与生产构建统一业务语义")
    capabilities = [
        ("人", "员工｜承包商｜访客｜组织", YELLOW),
        ("地", "实时定位｜轨迹｜区域进出｜三维地图", TEAL),
        ("票", "作业票｜合规｜流程｜工时", CYAN),
        ("警", "人员报警｜异常驻留｜疲劳风险", RED),
        ("设备", "健康度｜状态监测｜故障预测｜RUL", GREEN),
        ("生产", "电｜热｜气｜调度｜能效", "8B79E8"),
    ]
    for i, (name, desc, color) in enumerate(capabilities):
        col, row = i % 3, i // 3
        x, y = 0.65 + col * 4.17, 1.55 + row * 2.35
        add_box(slide, x, y, 3.72, 1.9, fill=PANEL, line=color)
        add_box(slide, x + 0.25, y + 0.26, 0.74, 0.74, fill=color, line=color, radius=True, width=0)
        add_text(slide, name, x + 0.25, y + 0.45, 0.74, 0.26, size=18, color=BG, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, name + "的管理", x + 1.23, y + 0.28, 2.1, 0.34, size=16, bold=True)
        add_text(slide, desc, x + 1.23, y + 0.83, 2.1, 0.6, size=10.5, color=MUTED)
    add_text(slide, "统一智能入口 · 60 项业务接口 · 动态工具路由 · 多模型协同", 0.75, 6.3, 11.8, 0.38, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "这页是能力总览。不要逐项念功能，强调六类对象可以互相关联，例如人员轨迹与作业票、设备报警与维护工单、生产负荷与安全风险。")

    # 06 架构
    slide = new_slide(prs, 6, "产品总体架构", "保留客户现有系统，以智能层形成统一入口和业务关联")
    layers = [
        ("交互层", "管理驾驶舱｜自然语言对话｜移动端｜报表与三维地图", YELLOW),
        ("智能层", "意图识别｜动态工具路由｜知识检索｜模型编排｜结果校验", TEAL),
        ("业务能力层", "人员安全｜作业合规｜能源运营｜预测性维护｜报警闭环", CYAN),
        ("数据接入层", "API｜数据库｜消息总线｜时序数据｜文件与制度知识库", GREEN),
        ("现场系统层", "定位｜门禁｜作业票｜DCS/SCADA｜EAM/CMMS｜IoT 终端", "667985"),
    ]
    for i, (name, desc, color) in enumerate(layers):
        y = 1.45 + i * 1.0
        add_box(slide, 0.85, y, 11.65, 0.72, fill=PANEL if i % 2 == 0 else PANEL_2, line=color)
        add_text(slide, name, 1.12, y + 0.2, 1.45, 0.25, size=13, color=color, bold=True)
        add_text(slide, desc, 2.75, y + 0.2, 9.25, 0.25, size=12, color=WHITE)
    add_text(slide, "核心原则：不推翻原有系统，不复制业务流程，用智能入口释放现有数据价值", 0.85, 6.62, 11.65, 0.32, size=14, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "客户担心改造量时重点讲：平台通过接口连接现有系统，不要求替换定位、门禁、作业票或设备管理平台。第一阶段从只读查询开始，降低上线风险。")

    # 07 人员
    slide = new_slide(prs, 7, "人员与承包商：从档案管理到现场态势", "回答三个问题：谁在现场、现在在哪、是否符合准入和作业要求")
    steps = [("身份", "员工/承包商/访客\n组织与岗位"), ("准入", "培训/资质/门禁\n考勤与排班"), ("在场", "实时位置/在线状态\n区域人数"), ("行为", "轨迹/进出/驻留\n劳动强度"), ("闭环", "报警/处置/复盘\n责任追溯")]
    for i, (title, desc) in enumerate(steps):
        x = 0.55 + i * 2.52
        add_box(slide, x, 1.65, 2.15, 1.55, fill=PANEL, line=TEAL if i < 4 else YELLOW)
        add_text(slide, title, x + 0.18, 1.88, 1.78, 0.3, size=16, color=TEAL if i < 4 else YELLOW, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.16, 2.34, 1.82, 0.55, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(slide, "›", x + 2.22, 2.15, 0.25, 0.35, size=24, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.72, 3.75, 12.0, 2.28, fill="0E1A22", line=LINE)
    add_text(slide, "客户可直接提问", 1.0, 4.05, 2.0, 0.3, size=14, color=YELLOW, bold=True)
    questions = ["“当前一号装置区有多少承包商？”", "“张伟今天经过了哪些区域？”", "“哪些人员存在异常驻留或疲劳风险？”"]
    add_rich_lines(slide, questions, 1.0, 4.55, 6.1, 1.2, size=13, color=WHITE, bullet=True, spacing=11)
    add_text(slide, "销售价值点", 8.0, 4.05, 1.6, 0.3, size=14, color=TEAL, bold=True)
    add_rich_lines(slide, ["缩短人员查找时间", "强化承包商全过程管理", "为应急处置提供位置依据"], 8.0, 4.55, 4.0, 1.2, size=13, color=WHITE, bullet=True, spacing=11)
    add_notes(slide, "人员定位不是终点。强调身份、准入、位置、行为、作业五类信息关联。客户已有定位系统也不冲突，我们提供的是统一分析和智能入口。")

    # 08 三维孪生
    slide = new_slide(prs, 8, "透明建筑数字孪生：让现场位置一眼可见", "以建筑透视方式呈现人员、风险区域、道路和空间状态")
    add_box(slide, 0.65, 1.45, 7.25, 5.2, fill="101920", line=LINE)
    # 厂区道路
    for y in (3.05, 5.2):
        road = add_box(slide, 0.9, y, 6.72, 0.28, fill="2B3439", line="2B3439", radius=False, width=0)
        add_text(slide, "-  -  -  -  -  -  -  -  -", 1.0, y + 0.04, 6.5, 0.16, size=8, color=YELLOW, align=PP_ALIGN.CENTER)
    buildings = [(1.05, 1.75, 1.65, 1.05, "中控楼"), (3.0, 1.68, 2.15, 1.15, "一号装置区"), (5.47, 1.82, 1.75, 0.98, "配电室"), (1.15, 3.65, 2.1, 1.05, "检修车间"), (3.65, 3.55, 1.55, 1.18, "罐区"), (5.55, 3.68, 1.55, 1.02, "动力泵房")]
    for i, (x, y, w, h, name) in enumerate(buildings):
        shape = add_box(slide, x, y, w, h, fill="243943", line=TEAL if i in (1, 5) else "66808D", radius=False, width=1.2)
        add_text(slide, name, x + 0.08, y + h - 0.25, w - 0.16, 0.18, size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    for x, y, name in [(3.62, 2.1, "张伟"), (4.25, 2.38, "李强"), (6.25, 4.05, "王磊")]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.19), Inches(0.19))
        circle.fill.solid(); circle.fill.fore_color.rgb = rgb(GREEN); circle.line.color.rgb = rgb(WHITE)
        add_text(slide, name, x - 0.23, y - 0.26, 0.66, 0.2, size=7.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "透明建筑 · 人员在建筑内部显示", 1.05, 5.72, 6.3, 0.32, size=12, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "客户价值", 8.35, 1.55, 2.0, 0.35, size=18, color=YELLOW, bold=True)
    add_rich_lines(slide, ["快速确认人员所在建筑和区域", "直观看到风险区内人员分布", "支持实时定位数据动态更新", "应急场景中辅助人员清点", "与人员名单和报警信息联动"], 8.35, 2.2, 4.05, 2.7, size=14, color=WHITE, bullet=True, spacing=14)
    add_badge(slide, "可对接 RFID / UWB / AoA / GPS", 8.35, 5.48, 3.45, TEAL)
    add_notes(slide, "展示时先问客户：发生报警时，多久能确认区域内有哪些人？然后演示自然语言提问和透明建筑内人员显示。不要承诺定位精度，精度取决于客户现场定位系统。")

    # 09 作业票
    slide = new_slide(prs, 9, "作业票合规：把票面要求与现场行为关联起来", "从单票查询升级为趋势、规则、流程与人员活动的综合分析")
    flow = [("票", "许可范围\n许可时间"), ("人", "责任人\n作业人员"), ("地", "实际区域\n进出轨迹"), ("规", "规则校验\n违规排行"), ("闭环", "处置结果\n复盘改进")]
    for i, (name, desc) in enumerate(flow):
        x = 0.65 + i * 2.5
        add_box(slide, x, 1.62, 1.86, 1.35, fill=PANEL, line=YELLOW if i in (0, 4) else TEAL)
        add_text(slide, name, x + 0.2, 1.84, 1.46, 0.32, size=18, color=YELLOW if i in (0, 4) else TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.18, 2.3, 1.5, 0.46, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4: add_text(slide, "→", x + 1.93, 2.05, 0.5, 0.3, size=18, color="667A84", align=PP_ALIGN.CENTER)
    items = [("合规统计", "按日期、票种、部门、单位聚合"), ("趋势分析", "按日/月查看合规变化"), ("规则排行", "识别高频违规规则"), ("流程分析", "定位审批与执行耗时"), ("工时核算", "按人员和承包商汇总")]
    for i, (title, desc) in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 0.72 + col * 4.05, 3.55 + row * 1.15
        add_box(slide, x, y, 3.65, 0.88, fill=PANEL_2, line=LINE)
        add_text(slide, title, x + 0.16, y + 0.14, 1.0, 0.25, size=11.5, color=TEAL, bold=True)
        add_text(slide, desc, x + 1.22, y + 0.14, 2.2, 0.42, size=10, color=MUTED)
    add_notes(slide, "如果客户已有电子作业票，重点不是替代，而是把票与人员定位、门禁和现场活动关联，回答‘票上允许什么、现场实际发生了什么’。")

    # 10 风险
    slide = new_slide(prs, 10, "报警与人员风险：从报警列表到处置优先级", "识别连续报警、高频区域、异常驻留和疲劳风险")
    add_metric(slide, "七类", "人员安全报警统一汇总", 0.72, 1.55, 2.7, RED)
    add_metric(slide, "连续", "识别重复与持续报警", 3.62, 1.55, 2.7, YELLOW)
    add_metric(slide, "时长", "统计发现到处置闭环", 6.52, 1.55, 2.7, TEAL)
    add_metric(slide, "强度", "驻留、距离与疲劳分析", 9.42, 1.55, 2.7, CYAN)
    add_box(slide, 0.72, 3.02, 12.0, 2.85, fill=PANEL, line=LINE)
    add_text(slide, "风险研判链路", 1.0, 3.3, 2.0, 0.32, size=16, color=YELLOW, bold=True)
    stages = [("发现", RED), ("定位", YELLOW), ("关联", CYAN), ("研判", TEAL), ("处置", GREEN), ("复盘", "8B79E8")]
    for i, (name, color) in enumerate(stages):
        x = 1.02 + i * 1.83
        add_box(slide, x, 4.0, 1.42, 0.66, fill=PANEL_2, line=color)
        add_text(slide, name, x + 0.08, 4.2, 1.26, 0.22, size=13, color=color, bold=True, align=PP_ALIGN.CENTER)
        if i < 5: add_text(slide, "→", x + 1.47, 4.18, 0.28, 0.22, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "输出：风险人员｜发生区域｜持续时间｜关联作业｜处置状态｜改进建议", 1.02, 5.1, 10.9, 0.32, size=13, color=WHITE, align=PP_ALIGN.CENTER)
    add_notes(slide, "强调平台输出优先级和关联信息，但不替代安全联锁或法定处置流程。高风险事件必须由规则和责任人员确认。")

    # 11 能源运营
    slide = new_slide(prs, 11, "电、热、气一体化运营", "一个智能入口覆盖生产运行、调度、能效和安全风险")
    domains = [
        ("电", "机组出力｜电网负荷｜变电站｜线路｜停电", YELLOW),
        ("热", "供回温｜压力｜流量｜热负荷｜热网报警", RED),
        ("气", "门站｜调压站｜供气量｜压力｜加臭状态", TEAL),
    ]
    for i, (name, desc, color) in enumerate(domains):
        x = 0.72 + i * 4.1
        add_box(slide, x, 1.55, 3.65, 2.0, fill=PANEL, line=color)
        add_text(slide, name, x + 0.22, 1.82, 0.7, 0.62, size=32, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 1.05, 1.85, 2.25, 0.9, size=12, color=WHITE)
    add_box(slide, 0.72, 4.1, 11.85, 1.62, fill=PANEL_2, line=LINE)
    add_text(slide, "综合管理视角", 1.02, 4.38, 1.9, 0.34, size=16, color=YELLOW, bold=True)
    add_rich_lines(slide, ["调度指令执行闭环", "煤耗、厂用电率、线损、热损和气损", "设备重载、导线温升、热网异常和能源安全风险"], 3.05, 4.34, 8.85, 1.1, size=13, color=WHITE, bullet=True, spacing=9)
    add_text(slide, "适合综合能源集团跨专业运营交流与管理驾驶舱", 0.75, 6.28, 11.8, 0.36, size=15, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "面向集团或综合能源客户时突出电热气协同。面向单一工厂时只讲客户相关专业，避免功能过多稀释价值。")

    # 12 预测维护
    slide = new_slide(prs, 12, "预测性维护：让设备风险早于故障被发现", "从状态采集、健康评价、风险预测到维护建议形成闭环")
    chain = [
        ("感知", "振动/温度/电流\n局放/油色谱", CYAN),
        ("评价", "健康度\n异常分数", TEAL),
        ("预测", "故障模式\n发生概率", YELLOW),
        ("决策", "剩余寿命\n建议窗口", RED),
        ("执行", "工单/备件\n停机计划", GREEN),
    ]
    for i, (title, desc, color) in enumerate(chain):
        x = 0.56 + i * 2.52
        add_box(slide, x, 1.55, 2.15, 1.45, fill=PANEL, line=color)
        add_text(slide, title, x + 0.15, 1.78, 1.85, 0.3, size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.15, 2.22, 1.85, 0.48, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4: add_text(slide, "→", x + 2.18, 2.06, 0.3, 0.25, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "标准监测终端", 0.75, 3.55, 2.0, 0.32, size=15, color=YELLOW, bold=True)
    terminals = ["无线振动", "无线测温", "电机电流", "局放在线", "油色谱", "压力流量", "红外测温", "边缘网关"]
    for i, terminal in enumerate(terminals):
        x = 0.75 + (i % 4) * 2.93
        y = 4.08 + (i // 4) * 0.75
        add_badge(slide, terminal, x, y, 2.45, TEAL if i < 4 else CYAN)
    add_box(slide, 0.75, 5.84, 11.75, 0.65, fill="2A2114", line=YELLOW)
    add_text(slide, "销售边界：预测结果用于辅助决策，正式应用需结合设备机理、历史故障与专业人员复核", 1.0, 6.05, 11.25, 0.24, size=11, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "不要承诺‘零故障’。正确表述是：帮助客户更早发现异常趋势、优化检修优先级和窗口。询问客户已有传感器、历史故障、检修记录和工单系统情况。")

    # 13 AI引擎
    slide = new_slide(prs, 13, "AI 引擎：不是把 60 个接口全部塞给模型", "动态工具路由让本地模型既能覆盖全部业务，又保持速度和稳定性")
    add_box(slide, 0.75, 1.5, 2.2, 1.0, fill=PANEL, line=YELLOW)
    add_text(slide, "自然语言问题", 0.92, 1.84, 1.86, 0.28, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 3.08, 1.78, 0.45, 0.35, size=22, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 3.62, 1.5, 2.25, 1.0, fill=PANEL, line=TEAL)
    add_text(slide, "业务域识别", 3.82, 1.84, 1.85, 0.28, size=16, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 6.02, 1.78, 0.45, 0.35, size=22, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 6.55, 1.5, 2.25, 1.0, fill=PANEL, line=CYAN)
    add_text(slide, "加载相关接口", 6.73, 1.84, 1.88, 0.28, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 8.95, 1.78, 0.45, 0.35, size=22, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 9.5, 1.5, 2.75, 1.0, fill=PANEL, line=GREEN)
    add_text(slide, "数据查询与结论", 9.7, 1.84, 2.35, 0.28, size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    examples = [
        ("人员在哪里？", "定位轨迹 + 三维地图", "8 个接口"),
        ("作业票是否合规？", "作业票 + 人员活动", "按需组合"),
        ("主变未来风险？", "输变配电 + 预测维护", "按需组合"),
        ("终端为什么离线？", "预测性维护", "8 个接口"),
    ]
    for i, (q, domain, count) in enumerate(examples):
        y = 3.15 + i * 0.78
        add_box(slide, 0.8, y, 11.65, 0.58, fill=PANEL_2, line=LINE)
        add_text(slide, q, 1.05, y + 0.16, 3.0, 0.22, size=11.5, color=WHITE, bold=True)
        add_text(slide, domain, 4.55, y + 0.16, 3.7, 0.22, size=11.5, color=TEAL)
        add_badge(slide, count, 9.9, y + 0.12, 1.65, YELLOW)
    add_notes(slide, "这是本地模型差异化重点。动态路由解决工具过多导致的上下文溢出、误选和响应慢，同时支持无关键词追问沿用上一业务域。")

    # 14 模型与安全
    slide = new_slide(prs, 14, "本地优先、多模型协同", "根据数据敏感度、任务复杂度和成本选择最合适的模型")
    columns = [
        ("本地模型", ["敏感数据不出本机", "实时查询与工具调用", "GPU 加速、按需启动", "动态路由覆盖全部接口"], TEAL),
        ("云端模型", ["复杂推理与长报告", "多模型灵活切换", "按调用量使用", "适合非敏感分析"], CYAN),
        ("治理与审计", ["用户与角色权限", "敏感字段最小展示", "提示词与工具调用记录", "可接入 Langfuse 评测"], YELLOW),
    ]
    for i, (title, bullets, color) in enumerate(columns):
        x = 0.65 + i * 4.18
        add_box(slide, x, 1.5, 3.75, 4.75, fill=PANEL, line=color)
        add_text(slide, title, x + 0.26, 1.85, 3.23, 0.4, size=20, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_rich_lines(slide, bullets, x + 0.34, 2.65, 3.05, 2.7, size=13, color=WHITE, bullet=True, spacing=15)
    add_notes(slide, "面对数据安全异议，优先讲本地模型和内网部署；面对效果异议，讲按任务切换模型。不要承诺绝对安全，正式项目仍需权限、审计、网络和主机安全体系。")

    # 15 场景闭环
    slide = new_slide(prs, 15, "典型场景：高风险作业全过程闭环", "将人员、位置、作业票和报警串成一条可追溯链路")
    stages = [
        ("作业前", "人员资质\n培训准入\n票面合规", TEAL),
        ("作业中", "人员到场\n区域范围\n停留与报警", YELLOW),
        ("作业后", "工时核算\n轨迹复盘\n违规闭环", GREEN),
    ]
    for i, (title, desc, color) in enumerate(stages):
        x = 0.85 + i * 4.15
        add_box(slide, x, 1.55, 3.45, 2.2, fill=PANEL, line=color)
        add_text(slide, title, x + 0.25, 1.88, 2.95, 0.36, size=19, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.25, 2.55, 2.95, 0.85, size=14, color=WHITE, align=PP_ALIGN.CENTER)
        if i < 2: add_text(slide, "→", x + 3.55, 2.45, 0.48, 0.38, size=25, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.85, 4.35, 11.65, 1.45, fill=PANEL_2, line=LINE)
    add_text(slide, "示例提问", 1.15, 4.68, 1.2, 0.3, size=14, color=YELLOW, bold=True)
    add_text(slide, "“今天受限空间作业区有哪些人员？对应作业票是否合规？是否发生越界或异常驻留？”", 2.55, 4.64, 9.25, 0.5, size=15, color=WHITE, bold=True)
    add_text(slide, "价值：把多系统取证时间压缩为一次查询，帮助安全管理人员快速聚焦异常", 1.15, 5.35, 10.7, 0.26, size=11.5, color=MUTED)
    add_notes(slide, "这是最适合安全部门的组合场景。演示时先查人员，再查票，再看地图和报警，最后让系统给出简短结论。")

    # 16 设备场景
    slide = new_slide(prs, 16, "典型场景：设备风险提前识别", "把实时监测数据转化为可执行的运维优先级")
    add_box(slide, 0.7, 1.5, 7.3, 4.9, fill=PANEL, line=LINE)
    add_text(slide, "设备健康趋势", 1.0, 1.82, 2.2, 0.3, size=16, color=TEAL, bold=True)
    # 简化趋势图
    for i in range(5):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.1), Inches(2.55 + i * 0.56), Inches(7.55), Inches(2.55 + i * 0.56))
        line.line.color.rgb = rgb("29404B"); line.line.width = Pt(0.7)
    points = [(1.2, 4.85), (2.0, 4.68), (2.8, 4.7), (3.6, 4.35), (4.4, 4.18), (5.2, 3.68), (6.0, 3.5), (6.8, 2.88), (7.45, 2.65)]
    for a, b in zip(points, points[1:]):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1]))
        line.line.color.rgb = rgb(YELLOW); line.line.width = Pt(2.4)
    add_text(slide, "预警阈值", 6.42, 2.25, 0.9, 0.2, size=8.5, color=RED)
    threshold = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.1), Inches(2.48), Inches(7.55), Inches(2.48))
    threshold.line.color.rgb = rgb(RED); threshold.line.dash_style = 2
    add_text(slide, "振动值持续上升", 1.1, 5.42, 2.2, 0.28, size=11, color=YELLOW, bold=True)
    add_text(slide, "预测故障模式：轴承磨损", 3.2, 5.42, 2.5, 0.28, size=11, color=WHITE)
    add_text(slide, "建议：7天内专项检查", 5.72, 5.42, 1.9, 0.28, size=11, color=TEAL)
    add_text(slide, "从数据到行动", 8.5, 1.65, 2.5, 0.36, size=18, color=YELLOW, bold=True)
    add_rich_lines(slide, ["识别异常趋势", "评估故障概率", "计算剩余寿命", "推荐检修窗口", "核查备件和工单"], 8.5, 2.35, 3.6, 2.8, size=15, color=WHITE, bullet=True, spacing=16)
    add_badge(slide, "辅助决策，不替代专业诊断", 8.5, 5.55, 3.05, RED)
    add_notes(slide, "设备部门更关心数据来源和算法可信度。先确认传感器与历史数据，再讨论模型。销售不要只展示风险概率，必须展示证据、趋势和建议动作。")

    # 17 价值
    slide = new_slide(prs, 17, "客户价值：从信息查询到运营闭环", "指标需在项目调研后共同确定，以下为价值评估方向")
    values = [
        ("更快", "跨系统查询与临时报表", "缩短信息获取时间", YELLOW),
        ("更准", "统一数据口径与关联分析", "减少人工判断遗漏", TEAL),
        ("更早", "异常趋势与风险预测", "提前安排处置窗口", CYAN),
        ("更稳", "本地部署与动态工具路由", "保障敏感场景可用性", GREEN),
    ]
    for i, (headline, scene, benefit, color) in enumerate(values):
        x = 0.65 + i * 3.1
        add_box(slide, x, 1.55, 2.75, 3.35, fill=PANEL, line=color)
        add_text(slide, headline, x + 0.18, 1.9, 2.38, 0.6, size=30, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, scene, x + 0.25, 2.85, 2.25, 0.62, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, benefit, x + 0.25, 3.75, 2.25, 0.62, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.75, 5.38, 11.75, 0.9, fill=PANEL_2, line=LINE)
    add_text(slide, "建议项目 KPI：查询响应时间｜报警闭环时长｜作业票合规率｜设备异常提前量｜非计划停机时长", 1.02, 5.69, 11.2, 0.3, size=13, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "不要在没有基线数据时承诺具体百分比。销售应与客户共同定义现状基线、目标指标和验收口径。")

    # 18 需求访谈
    slide = new_slide(prs, 18, "销售需求访谈：必须问清的 12 个问题", "通过问题判断客户成熟度、项目边界和切入场景")
    questions = [
        "目前有哪些人员定位、门禁和作业票系统？",
        "发生报警后，确认人员位置通常需要多久？",
        "承包商准入、作业和工时由哪些部门管理？",
        "作业票合规主要依赖人工还是系统规则？",
        "哪些高风险区域最需要实时人员态势？",
        "现有定位技术和实际覆盖范围是什么？",
        "重点设备已经采集哪些状态量？",
        "是否具备历史故障、缺陷和检修工单数据？",
        "最关注哪些非计划停机或设备故障？",
        "数据是否允许上云？内网模型有哪些要求？",
        "现有系统能否提供 API、数据库或消息接口？",
        "谁负责验收？希望改善哪些量化指标？",
    ]
    for i, question in enumerate(questions):
        col, row = i % 2, i // 2
        x, y = 0.72 + col * 6.17, 1.42 + row * 0.86
        add_box(slide, x, y, 5.8, 0.62, fill=PANEL if row % 2 == 0 else PANEL_2, line=LINE)
        add_text(slide, f"{i+1:02d}", x + 0.14, y + 0.18, 0.4, 0.22, size=10, color=YELLOW, bold=True)
        add_text(slide, question, x + 0.63, y + 0.16, 4.9, 0.3, size=10.5, color=WHITE)
    add_notes(slide, "需求访谈的目标不是马上报价，而是确认数据基础、关键痛点、决策链和验收指标。至少拿到一个可量化场景再进入方案阶段。")

    # 19 标准演示
    slide = new_slide(prs, 19, "标准产品演示路径｜15 分钟", "先给结论，再展示关联，再讲技术；不要从配置页面开始")
    demo = [
        ("01", "2分钟", "产品定位", "一句话说明统一智能入口与六位一体能力"),
        ("02", "3分钟", "人员定位", "查询人员位置并在透明三维建筑中显示"),
        ("03", "3分钟", "作业合规", "查看作业票、合规结果和违规规则"),
        ("04", "3分钟", "预测维护", "查询离线终端、高风险设备和维护建议"),
        ("05", "2分钟", "本地模型", "展示动态工具路由、流式输出和硬件资源"),
        ("06", "2分钟", "价值收口", "回到客户痛点，确认下一步数据接口调研"),
    ]
    for i, (num, duration, title, desc) in enumerate(demo):
        y = 1.4 + i * 0.86
        add_text(slide, num, 0.85, y + 0.16, 0.45, 0.25, size=11, color=YELLOW, bold=True)
        add_badge(slide, duration, 1.48, y + 0.1, 0.9, TEAL)
        add_text(slide, title, 2.75, y + 0.13, 1.4, 0.28, size=13, color=WHITE, bold=True)
        add_box(slide, 4.25, y, 8.0, 0.58, fill=PANEL, line=LINE)
        add_text(slide, desc, 4.52, y + 0.16, 7.45, 0.25, size=11, color=MUTED)
    add_notes(slide, "演示过程中不要临时尝试未验证问题。使用标准问题集；每个场景都遵循‘问题—接口调用—可视化—一句话结论—客户价值’结构。")

    # 20 异议处理
    slide = new_slide(prs, 20, "常见客户异议与标准回应", "回应原则：承认边界、说明机制、给出验证方法")
    objections = [
        ("“我们已经有定位系统了。”", "平台不替代定位系统，而是关联人员、作业票、门禁、报警和三维态势，释放已有数据价值。"),
        ("“大模型会不会编造数据？”", "涉及业务数据必须调用接口；无数据明确返回查不到，并通过工具日志和评测体系持续校验。"),
        ("“敏感数据不能上云。”", "支持本地模型与内网部署，敏感字段最小展示；云端模型可按企业策略关闭。"),
        ("“预测性维护可靠吗？”", "提供趋势、证据、置信度和建议窗口，定位为辅助决策；通过历史数据回测和现场试点验证。"),
        ("“接口太多，本地模型能处理吗？”", "动态工具路由覆盖全部业务接口，每次只加载相关工具组，降低上下文和误选风险。"),
    ]
    for i, (question, answer) in enumerate(objections):
        y = 1.4 + i * 1.05
        add_box(slide, 0.72, y, 4.0, 0.8, fill="21191A", line=RED)
        add_text(slide, question, 0.96, y + 0.2, 3.5, 0.36, size=11.5, color="FFB0B0", bold=True)
        add_box(slide, 4.95, y, 7.62, 0.8, fill=PANEL, line=TEAL)
        add_text(slide, answer, 5.2, y + 0.15, 7.1, 0.48, size=10.5, color=WHITE)
    add_notes(slide, "不要与客户争辩。先确认顾虑合理，再说明产品机制和验证方式。对数据安全、算法效果和定位精度等问题必须保留实施边界。")

    # 21 交付路径
    slide = new_slide(prs, 21, "产品落地路径", "从一个高价值场景切入，逐步扩展数据范围和业务闭环")
    phases = [
        ("01", "场景与数据诊断", "业务访谈｜接口盘点｜指标基线", "2–4周", YELLOW),
        ("02", "重点场景上线", "人员定位｜作业合规｜重点设备", "4–8周", TEAL),
        ("03", "系统与流程贯通", "权限审计｜工单闭环｜多系统关联", "按范围", CYAN),
        ("04", "规模化运营", "多厂区复制｜模型评测｜持续优化", "持续", GREEN),
    ]
    for i, (num, title, desc, duration, color) in enumerate(phases):
        x = 0.6 + i * 3.16
        add_box(slide, x, 1.55, 2.8, 3.85, fill=PANEL, line=color)
        add_text(slide, num, x + 0.22, 1.85, 0.55, 0.32, size=14, color=color, bold=True)
        add_text(slide, title, x + 0.22, 2.42, 2.35, 0.58, size=17, color=WHITE, bold=True)
        add_text(slide, desc, x + 0.22, 3.35, 2.35, 0.8, size=12, color=MUTED)
        add_badge(slide, duration, x + 0.22, 4.72, 1.1, color)
    add_text(slide, "推荐切入：高风险作业人员态势 或 重点设备预测性维护", 0.72, 6.05, 11.85, 0.4, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "落地策略是小场景、真实数据、可量化价值。第一阶段不追求覆盖全部系统，优先选择数据基础好、管理痛点强、责任部门明确的场景。")

    # 22 收尾
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_box(slide, 0, 0, 0.16, H, fill=YELLOW, line=YELLOW, radius=False, width=0)
    add_text(slide, "人员安全智能运营中心", 0.85, 1.2, 8.0, 0.62, size=32, bold=True)
    add_text(slide, "让人员安全看得见，让设备风险早发现", 0.88, 2.15, 9.3, 0.55, size=23, color=YELLOW, bold=True)
    add_text(slide, "连接人、地、票、警、设备与生产\n以可信数据驱动工业现场安全运营升级", 0.9, 3.15, 7.6, 1.1, size=18, color="C1D0D6")
    add_box(slide, 9.35, 1.25, 2.6, 2.6, fill=PANEL, line=TEAL)
    add_text(slide, "下一步", 9.83, 1.7, 1.65, 0.38, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "选择一个场景\n完成一次数据诊断\n验证一项业务价值", 9.68, 2.35, 1.95, 1.0, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_badge(slide, "工业智能 · 安全运营 · 预测维护", 0.9, 5.45, 3.55, YELLOW)
    add_text(slide, "THANK YOU", 0.92, 6.65, 2.0, 0.28, size=10, color="607985")
    add_notes(slide, "收尾不要停在‘谢谢’。邀请客户共同选择一个场景，安排业务、设备和信息化人员参加接口与数据诊断会。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT, len(prs.slides)


if __name__ == "__main__":
    output, count = build_deck()
    print(f"generated={output}")
    print(f"slides={count}")

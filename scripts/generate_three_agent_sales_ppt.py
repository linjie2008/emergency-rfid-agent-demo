"""生成三智能体售前与销售培训 PPT。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "三智能体智能运营中心-售前销售培训.pptx"
W, H = 13.333, 7.5
BG, PANEL, PANEL2 = "07131D", "102532", "16313F"
WHITE, MUTED, LINE = "F4F8FA", "9DB2BD", "294757"
TEAL, BLUE, ORANGE, GREEN, RED = "39D2C0", "52A9E8", "F2AE3D", "55C98D", "EF6A6A"
FONT = "Microsoft YaHei"


def rgb(c): return RGBColor.from_string(c)


def add_box(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True, width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(width)
    return shape


def add_text(slide, text, x, y, w, h, size=16, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=.03):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size)
    r.font.bold = bold; r.font.color.rgb = rgb(color)
    return shape


def add_lines(slide, lines, x, y, w, h, size=13, color=WHITE, bullet=True, gap=8):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.04)
    for i, item in enumerate(lines):
        if isinstance(item, tuple): text, item_color, bold = item
        else: text, item_color, bold = item, color, False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("•  " if bullet else "") + text
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = rgb(item_color); p.space_after = Pt(gap)
    return shape


def add_title(slide, no, title, subtitle=""):
    add_box(slide, .45, .31, .08, .5, TEAL, TEAL, False, 0)
    add_text(slide, title, .68, .25, 10.9, .48, 25, WHITE, True)
    if subtitle: add_text(slide, subtitle, .7, .77, 11.5, .28, 10.5, MUTED)
    add_text(slide, f"{no:02d}", 12.18, .3, .55, .3, 11, TEAL, True, PP_ALIGN.RIGHT)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(.5), Inches(1.13), Inches(12.82), Inches(1.13))
    ln.line.color.rgb = rgb(LINE)


def add_footer(slide, no):
    add_text(slide, "三智能体智能运营中心｜售前与销售培训", .52, 7.16, 5.2, .16, 8, "65808D")
    add_text(slide, f"{no:02d}", 12.2, 7.14, .55, .17, 8, "65808D", align=PP_ALIGN.RIGHT)


def notes(slide, text):
    try: slide.notes_slide.notes_text_frame.text = text
    except Exception: pass


def new_slide(prs, no, title, subtitle=""):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    add_title(s, no, title, subtitle); add_footer(s, no); return s


def card(slide, x, y, w, h, title, body, color=TEAL, tag=""):
    add_box(slide, x, y, w, h, PANEL, LINE)
    add_box(slide, x+.16, y+.17, .07, .38, color, color, False, 0)
    add_text(slide, title, x+.34, y+.15, w-.5, .34, 15, WHITE, True)
    add_text(slide, body, x+.18, y+.7, w-.36, h-.82, 10.5, MUTED)
    if tag:
        add_box(slide, x+w-1.22, y+.16, 1.03, .3, PANEL2, color)
        add_text(slide, tag, x+w-1.18, y+.205, .95, .17, 8.5, color, True, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, color=TEAL):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(2.2); ln.line.end_arrowhead = True


def build():
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    prs.core_properties.title = "三智能体智能运营中心｜售前与销售培训"
    prs.core_properties.subject = "医疗、人员安全、政企楼宇三智能体销售方法与演示架构"
    prs.core_properties.author = "智能运营中心"

    # 01 封面
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    for i, c in enumerate((BLUE, ORANGE, TEAL)):
        add_box(s, 8.8+i*.72, -.3+i*.55, 4.3, 1.55, c, c, True, 0).rotation = -28
    add_text(s, "三智能体智能运营中心", .72, 1.7, 9.2, .75, 36, WHITE, True)
    add_text(s, "医疗 × 人员安全 × 政企楼宇", .76, 2.62, 8, .45, 21, TEAL, True)
    add_text(s, "售前与销售培训", .76, 3.28, 5.3, .45, 19, ORANGE, True)
    add_text(s, "从客户痛点、产品价值到演示成交的一套统一话术", .78, 4.05, 8.6, .36, 14, MUTED)
    for i, (sname, c) in enumerate((("医疗运营", BLUE), ("人员安全", ORANGE), ("政企楼宇", TEAL))):
        add_box(s, .78+i*2.05, 5.25, 1.78, .48, PANEL2, c)
        add_text(s, sname, .82+i*2.05, 5.37, 1.7, .21, 10.5, c, True, PP_ALIGN.CENTER)
    add_text(s, "PRE-SALES & SALES ENABLEMENT", .79, 6.74, 5, .2, 8.5, "607C89")
    notes(s, "开场先定调：这不是三个孤立的聊天机器人，而是一套可复用的智能运营平台，面向三个行业场景使用独立数据和接口。")

    # 02 培训目标
    s = new_slide(prs, 2, "培训结束后，销售要能做到四件事", "不要求讲深技术，但必须把价值、边界和演示路径讲清楚")
    items = [("讲得清", "用一分钟说清三智能体是什么、为什么不是普通聊天机器人", BLUE),
             ("找得准", "识别客户角色、业务痛点、现有系统和可落地切入口", ORANGE),
             ("演得稳", "用固定问题展示接口调用、数据隔离、定位与可视化", TEAL),
             ("推得动", "从演示推进到接口调研、试点范围、预算与实施计划", GREEN)]
    for i,(t,b,c) in enumerate(items):
        x=.65+(i%2)*6.15; y=1.55+(i//2)*2.45
        card(s,x,y,5.72,1.9,t,b,c,f"0{i+1}")
    notes(s, "培训不是让销售背功能表，而是形成完整动作：发现痛点—选择智能体—演示证据—推动下一步。")

    # 03 产品组合
    s = new_slide(prs, 3, "一套平台，三个行业智能体", "交互体验一致，业务接口、数据与会话记忆按场景隔离")
    agents = [("医疗智能体", "患者/医护定位\n急诊绿通与超时\n医疗资产与能效", BLUE),
              ("人员安全智能体", "人员/资产定位\n考勤门禁与作业票\n报警与预测维护", ORANGE),
              ("政企楼宇智能体", "人员/资产定位\n能源空调与机电\n安防消防与维保", TEAL)]
    for i,(t,b,c) in enumerate(agents):
        x=.62+i*4.2; card(s,x,1.55,3.78,3.55,t,b,c,"独立场景")
        add_text(s, "自然语言提问 → 对应业务接口 → 可核验结果", x+.2, 4.45, 3.38, .36, 10, WHITE, True, PP_ALIGN.CENTER)
    add_box(s, 1.45, 5.55, 10.4, .72, PANEL2, LINE)
    add_text(s, "共同底座：统一交互｜工具编排｜模型切换｜权限审计｜图表与定位可视化", 1.7, 5.78, 9.9, .25, 14, TEAL, True, PP_ALIGN.CENTER)
    notes(s, "销售表达：客户买的是统一智能运营能力，首期可只落一个场景，后续复用平台扩展其他场景。")

    # 04 为什么现在需要
    s = new_slide(prs, 4, "客户不是缺系统，而是缺少统一决策入口", "从“系统建设”转向“把已有系统用起来”")
    pains = [("系统多", "不同系统反复登录，查询链路长"), ("数据散", "人员、设备、事件难以关联"),
             ("告警多", "只报异常，缺少优先级和处置建议"), ("汇报慢", "临时报表靠人工导出拼接"),
             ("使用难", "一线人员记不住菜单和字段"), ("数据敏感", "核心业务不能无边界上云")]
    for i,(t,b) in enumerate(pains):
        x=.62+(i%3)*4.18; y=1.5+(i//3)*2.35
        card(s,x,y,3.76,1.78,t,b,RED if i<3 else ORANGE,f"痛点{i+1}")
    add_text(s, "销售切入句：保留客户现有系统，用一个可以直接提问的智能入口连接它们。", .75, 6.18, 11.8, .34, 16, ORANGE, True, PP_ALIGN.CENTER)
    notes(s, "不要一上来讲大模型参数。先问客户跨系统查一个问题需要多久、谁来整理、是否能追溯来源。")

    # 05 通俗工作原理
    s = new_slide(prs, 5, "智能体怎么工作：像一个懂业务的数字值班员", "听懂问题、找到系统、调用接口、组织答案，四步完成")
    steps = [("1", "听懂问题", "判断用户是在问位置、告警、能耗还是流程", BLUE),
             ("2", "选择工具", "从允许的业务接口中选择正确查询能力", ORANGE),
             ("3", "读取数据", "调用客户现有 API/数据库，获得实时可信结果", TEAL),
             ("4", "组织答案", "输出结论、表格、图表、位置图或原理图", GREEN)]
    for i,(n,t,b,c) in enumerate(steps):
        x=.48+i*3.22
        add_box(s,x,1.85,2.65,3.15,PANEL, c)
        add_box(s,x+.86,1.48,.9,.9,c,c,True,0)
        add_text(s,n,x+.86,1.69,.9,.28,20,BG,True,PP_ALIGN.CENTER)
        add_text(s,t,x+.22,2.25,2.21,.38,17,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,b,x+.3,3.08,2.05,1.1,11,MUTED,False,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
        if i<3: arrow(s,x+2.68,3.35,x+3.16,3.35,c)
    add_text(s, "关键区别：普通聊天机器人主要“生成文字”；业务智能体必须“调用工具、依据数据、留下过程”。", .75, 5.78, 11.85, .44, 15, ORANGE, True, PP_ALIGN.CENTER)
    notes(s, "类比数字值班员：他不凭记忆报数字，而是去正确系统查数据，再用管理者听得懂的方式回答。")

    # 06 架构
    s = new_slide(prs, 6, "售前需要掌握的总体架构", "从上到下讲：入口、智能、业务、接口、现场系统")
    layers = [("交互入口", "PC 网页｜Telegram｜驾驶舱｜移动端", BLUE),
              ("智能编排", "意图识别｜场景路由｜工具调用｜结果组织", ORANGE),
              ("三类智能体", "医疗｜人员安全｜政企楼宇", TEAL),
              ("数据接口", "API｜数据库｜消息总线｜文件知识库", GREEN),
              ("客户系统", "定位｜门禁｜业务系统｜IoT｜能耗与设备监测", "718A97")]
    for i,(t,b,c) in enumerate(layers):
        y=1.38+i*1.02
        add_box(s,1.05,y,11.25,.72,PANEL if i%2==0 else PANEL2,c)
        add_text(s,t,1.35,y+.19,1.55,.25,13,c,True)
        add_text(s,b,3.2,y+.18,8.6,.25,13,WHITE)
        if i<4: arrow(s,6.66,y+.73,6.66,y+.98,c)
    notes(s, "客户问是否替换现有系统：明确回答不替换。智能体通过标准接口调用现有能力，降低重复建设。")

    # 07 隔离
    s = new_slide(prs, 7, "三个智能体为什么要隔离", "同一平台不等于混用数据：场景工具、业务数据、会话记忆分别隔离")
    cols=[("医疗命名空间", "患者、医护、设备\n医疗接口白名单\n医疗独立会话", BLUE),
          ("安全命名空间", "员工、承包商、资产\n安全接口白名单\n安全独立会话", ORANGE),
          ("楼宇命名空间", "员工、访客、设施\n楼宇接口白名单\n楼宇独立会话", TEAL)]
    for i,(t,b,c) in enumerate(cols): card(s,.65+i*4.18,1.65,3.72,3.65,t,b,c,"数据隔离")
    add_text(s, "共享的只有通用能力：登录入口、模型管理、表格/图表渲染和运维框架。", .75,5.75,11.8,.36,14,WHITE,True,PP_ALIGN.CENTER)
    notes(s, "强调隔离价值：减少串场、误调用和数据越权；正式项目还可叠加用户、角色、字段级权限与审计。")

    # 08 医疗
    s = new_slide(prs, 8, "医疗智能体：让急诊流程与医疗资产看得见", "核心购买者：急诊科、医务处、设备处、信息中心")
    card(s,.65,1.55,3.7,4.65,"客户痛点","患者流转不透明\n关键节点超时难发现\n设备找不到、利用率不清\n跨系统查询耗时",RED,"为什么买")
    card(s,4.78,1.55,3.7,4.65,"核心能力","患者与医护定位\n急诊绿通轨迹与停留\n超时预警与通道统计\n医疗资产定位、报警、能效",BLUE,"卖什么")
    card(s,8.91,1.55,3.7,4.65,"销售价值","提升急诊协同效率\n缩短查找患者/设备时间\n为流程优化提供数据证据\n降低资产闲置与丢失风险",GREEN,"带来什么")
    notes(s, "不要宣称替代临床决策。定位为运营与流程协同工具，医疗诊疗仍由专业人员负责。")

    # 09 医疗演示
    s = new_slide(prs, 9, "医疗智能体演示脚本", "三问建立完整故事：患者在哪里 → 是否超时 → 设备在哪里")
    qs=[("01 患者定位", "“张三现在在哪里？今天经过哪些区域？”", "实时位置 + 轨迹"),
        ("02 流程预警", "“当前抢救室有哪些停留超时患者？”", "超时清单 + 结论"),
        ("03 资产运营", "“输液泵 12 号在哪里？当前设备告警如何？”", "资产定位 + 告警")]
    for i,(t,q,o) in enumerate(qs):
        y=1.55+i*1.55; add_box(s,.75,y,11.8,1.15,PANEL,BLUE)
        add_text(s,t,1.0,y+.2,1.75,.28,13,BLUE,True)
        add_text(s,q,2.95,y+.18,6.1,.34,14,WHITE,True)
        add_text(s,o,9.45,y+.2,2.65,.28,11,GREEN,True,PP_ALIGN.CENTER)
    add_text(s,"演示收口：从“找人找设备”扩展到“流程效率与资产运营”。",.8,6.32,11.7,.3,15,ORANGE,True,PP_ALIGN.CENTER)
    notes(s, "每个问题先让客户猜需要打开几个系统，再展示一句自然语言得到结果。")

    # 10 人员安全
    s = new_slide(prs, 10, "人员安全智能体：围绕人、地、票、警形成闭环", "核心购买者：安全管理部、生产运营部、设备部、数字化部门")
    card(s,.65,1.55,3.7,4.65,"客户痛点","人员位置难确认\n承包商过程难监管\n作业票与轨迹割裂\n报警多、处置优先级不清",RED,"为什么买")
    card(s,4.78,1.55,3.7,4.65,"核心能力","人员与生产资产定位\n电子围栏、考勤与门禁\n作业票合规和有效工时\n报警分析与预测性维护",ORANGE,"卖什么")
    card(s,8.91,1.55,3.7,4.65,"销售价值","提升应急找人效率\n让高风险作业可追溯\n识别承包商与疲劳风险\n推动设备风险前置管理",GREEN,"带来什么")
    notes(s, "核心业务是人与资产定位，但销售不要只卖定位硬件，要卖定位数据与作业、安全、设备运营的关联价值。")

    # 11 安全演示
    s = new_slide(prs, 11, "人员安全智能体演示脚本", "从单点查询进入业务闭环，再展示可视化")
    qs=[("01 找人", "“张伟现在在哪里？今天有哪些区域停留？”", "定位 + 轨迹"),
        ("02 查作业", "“今天高风险作业票有哪些，合规情况如何？”", "作业票 + 合规"),
        ("03 看风险", "“当前有哪些未处理报警？哪些人员需优先关注？”", "报警 + 排序"),
        ("04 做展示", "“生成报警柱状图 / 画人员安全原理图。”", "图表 + ComfyUI")]
    for i,(t,q,o) in enumerate(qs):
        y=1.38+i*1.25; add_box(s,.7,y,11.9,.92,PANEL,ORANGE)
        add_text(s,t,.95,y+.17,1.65,.25,12,ORANGE,True)
        add_text(s,q,2.75,y+.15,6.65,.3,13,WHITE,True)
        add_text(s,o,9.7,y+.17,2.45,.24,11,GREEN,True,PP_ALIGN.CENTER)
    add_text(s,"触发词要明确：柱状图/折线图/饼图调用图表；原理图/示意图/架构图调用 ComfyUI。",.78,6.42,11.75,.3,12,MUTED,True,PP_ALIGN.CENTER)
    notes(s, "演示顺序建议固定，避免随机问答。定位—作业票—报警—可视化，四步最容易形成完整价值感。")

    # 12 楼宇
    s = new_slide(prs, 12, "政企楼宇智能体：统一机电、能源、安防与服务运营", "核心购买者：行政后勤、物业公司、能源管理、信息中心")
    card(s,.65,1.55,3.7,4.65,"客户痛点","楼宇子系统数量多\n能耗与设备状态分散\n告警难统一排序\n巡检维保依赖人工经验",RED,"为什么买")
    card(s,4.78,1.55,3.7,4.65,"核心能力","人员与设施资产定位\n电力、空调、水务、照明\n停车、电梯、安防、消防\n告警研判与维护台账",TEAL,"卖什么")
    card(s,8.91,1.55,3.7,4.65,"销售价值","统一运营视图\n提高能源与设施管理效率\n缩短告警研判时间\n支撑物业服务与节能优化",GREEN,"带来什么")
    notes(s, "楼宇数据模型依据调研清单构造。演示数据用于说明接口能力，正式项目需要完成真实系统接口映射。")

    # 13 楼宇演示
    s = new_slide(prs, 13, "政企楼宇智能体演示脚本", "先看全局，再下钻到系统、告警和处置")
    qs=[("01 综合态势", "“汇总总部大楼当前综合运营态势。”", "关键指标 + 异常"),
        ("02 能源空调", "“当前能耗和空调系统有哪些异常？”", "能耗 + 设备状态"),
        ("03 安防消防", "“列出未处理告警并给出处置优先级。”", "告警 + 研判"),
        ("04 人与资产", "“楼宇人员和重点设施资产分别在哪里？”", "位置 + 状态")]
    for i,(t,q,o) in enumerate(qs):
        y=1.38+i*1.25; add_box(s,.7,y,11.9,.92,PANEL,TEAL)
        add_text(s,t,.95,y+.17,1.75,.25,12,TEAL,True)
        add_text(s,q,2.85,y+.15,6.55,.3,13,WHITE,True)
        add_text(s,o,9.7,y+.17,2.45,.24,11,GREEN,True,PP_ALIGN.CENTER)
    add_text(s,"演示收口：从传统楼控“看数据”，升级为运营人员“问问题、得结论、看建议”。",.78,6.42,11.75,.3,13,ORANGE,True,PP_ALIGN.CENTER)
    notes(s, "优先选择客户熟悉的楼宇设备或告警问题，让客户看到智能体不是做大屏，而是降低运营分析门槛。")

    # 14 核心定位
    s = new_slide(prs, 14, "共同核心：人与资产定位是三类场景的时空底座", "接口契约一致，场景数据隔离，业务解释各不相同")
    add_box(s,4.82,2.22,3.7,1.65,PANEL2,ORANGE)
    add_text(s,"统一定位能力",5.17,2.53,3,.36,22,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"实时位置｜历史轨迹｜区域进出",5.1,3.15,3.15,.24,11,MUTED,False,PP_ALIGN.CENTER)
    endpoints=[("医院", "患者、医护、医疗设备", BLUE, 1.1,1.4), ("工业园区", "员工、承包商、生产资产", ORANGE,1.1,4.55), ("政企楼宇", "员工、访客、设施资产", TEAL,9.55,3.0)]
    for t,b,c,x,y in endpoints:
        card(s,x,y,2.8,1.35,t,b,c)
        arrow(s, x+2.8 if x<4 else x, y+.67, 4.75 if x<4 else 8.58, 3.05, c)
    add_text(s,"销售升级路径：定位项目 → 轨迹与区域规则 → 告警联动 → 业务流程智能化",1.15,6.25,11,.32,15,GREEN,True,PP_ALIGN.CENTER)
    notes(s, "这是公司核心业务页。把定位作为数据底座和切入口，但强调后续价值来自与医疗流程、作业票、楼宇运营的深度结合。")

    # 15 销售架构
    s = new_slide(prs, 15, "销售培训架构：从产品介绍到成交推进", "每次客户交流都围绕五个动作展开")
    stages=[("1 选场景","先确认客户属于医院、工业园区还是政企楼宇",BLUE),
            ("2 找痛点","定位查找、跨系统查询、告警研判、报表效率",RED),
            ("3 演证据","让智能体调用接口，展示数据来源与工具过程",ORANGE),
            ("4 定试点","确定一个组织、区域、系统和三到五个问题",TEAL),
            ("5 推商用","接口调研、权限、安全、部署、验收指标",GREEN)]
    for i,(t,b,c) in enumerate(stages):
        x=.42+i*2.57; add_box(s,x,1.72,2.18,3.75,PANEL,c)
        add_text(s,t,x+.16,2.03,1.86,.45,15,c,True,PP_ALIGN.CENTER)
        add_text(s,b,x+.25,2.98,1.68,1.6,11,MUTED,False,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
        if i<4: arrow(s,x+2.2,3.58,x+2.52,3.58,c)
    add_text(s,"目标不是“客户觉得很智能”，而是“客户同意提供接口清单并进入试点”。",.8,6.15,11.7,.38,16,ORANGE,True,PP_ALIGN.CENTER)
    notes(s, "每次演示结束必须约定下一步材料：系统清单、接口负责人、试点区域、验收问题集。")

    # 16 商机发现
    s = new_slide(prs, 16, "售前调研：六个问题快速判断商机质量", "答案越具体，越容易定义可成交的试点范围")
    questions=[("现有系统", "现在有哪些定位、门禁、业务和设备系统？"), ("高频查询", "管理人员每天最常查的三个问题是什么？"),
               ("协同成本", "一次跨系统查询或临时汇报需要多少人、多久？"), ("数据条件", "是否有 API、数据库或消息接口可用？"),
               ("安全要求", "数据能否上云？是否必须本地部署？"), ("成功标准", "客户希望试点后哪个指标发生变化？")]
    for i,(t,b) in enumerate(questions):
        x=.62+(i%2)*6.15; y=1.38+(i//2)*1.68
        card(s,x,y,5.72,1.25,t,b,BLUE if i%2==0 else TEAL,f"Q{i+1}")
    notes(s, "销售至少拿到：系统清单、问题清单、数据负责人、部署边界、决策链和时间计划。")

    # 17 演示规范
    s = new_slide(prs, 17, "演示成功的关键：固定故事线，不做无边界问答", "先准备问题，再准备接口数据和异常案例")
    good=["每个智能体准备 3—5 个黄金问题", "先展示结论，再展开表格、图表和轨迹", "主动说明数据来源与演示数据边界", "切换场景时强调数据和接口隔离", "生成图表时明确说柱状图/折线图/饼图"]
    bad=["让客户随意测试完全未知问题", "把所有能力一次性堆在一页", "把模拟数据说成客户真实数据", "只展示聊天，不展示工具调用", "承诺未验证的接口、准确率和自动处置"]
    card(s,.72,1.55,5.72,4.85,"推荐动作","\n".join("✓ "+x for x in good),GREEN,"DO")
    card(s,6.88,1.55,5.72,4.85,"避免动作","\n".join("× "+x for x in bad),RED,"DON'T")
    notes(s, "演示的稳定性高于花哨程度。每一个展示结果都应能回答：数据从哪里来、为什么可信、下一步能做什么。")

    # 18 异议处理
    s = new_slide(prs, 18, "常见异议与标准回答", "回答原则：先承认合理顾虑，再给出产品机制和落地路径")
    rows=[("“不就是聊天机器人吗？”", "不是。对话只是入口，核心是业务工具、接口调用、数据隔离与可追溯结果。"),
          ("“我们已经有很多系统。”", "不替换现有系统，通过接口形成统一查询和跨系统关联。"),
          ("“数据不能上云。”", "支持本地模型和内网部署，敏感场景可不出企业网络。"),
          ("“模型会不会胡说？”", "数字必须来自工具返回；无数据时明确说明，关键结果可保留调用过程。"),
          ("“三个场景会不会串数据？”", "每个智能体使用独立工具白名单、数据命名空间和会话记忆。")]
    for i,(q,a) in enumerate(rows):
        y=1.35+i*1.06; add_box(s,.65,y,12.0,.82,PANEL,LINE)
        add_text(s,q,.88,y+.16,3.1,.28,12,ORANGE,True)
        add_text(s,a,4.15,y+.14,8.05,.36,11.5,WHITE)
    notes(s, "不要用绝对化承诺，例如零幻觉、百分之百准确、无需实施。强调机制、边界和可验证性。")

    # 19 部署
    s = new_slide(prs, 19, "部署与商务组合", "根据数据敏感度、复杂度和预算选择合适方案")
    opts=[("演示验证", "模拟数据\n快速展示\n验证问题与流程", BLUE, "1—2 周"),
          ("单场景试点", "接入核心接口\n一个区域/科室\n形成验收问题集", ORANGE, "4—8 周"),
          ("内网商用", "本地模型\n权限与审计\n高可用与运维", TEAL, "分阶段"),
          ("多场景扩展", "复用统一平台\n新增数据适配\n跨组织推广", GREEN, "持续建设")]
    for i,(t,b,c,tag) in enumerate(opts): card(s,.55+i*3.18,1.55,2.82,4.55,t,b,c,tag)
    add_text(s,"商务拆分建议：平台软件 + 场景智能体 + 数据接口适配 + 定位硬件/实施 + 运维服务",.72,6.35,11.9,.28,13,WHITE,True,PP_ALIGN.CENTER)
    notes(s, "具体周期与报价以接口数量、数据质量、权限体系、部署环境和定制范围评估为准。")

    # 20 试点验收
    s = new_slide(prs, 20, "从 Demo 到商用：用可验收问题驱动实施", "避免泛化目标，把成功定义为一组能稳定回答的业务问题")
    stages=[("接口梳理", "系统、字段、频率、权限、负责人", BLUE), ("数据映射", "人员/资产主数据、区域、事件口径", ORANGE),
            ("问题集建设", "选择 20—50 个高价值业务问题", TEAL), ("联调评测", "工具选择、答案准确、延迟、权限", GREEN),
            ("试运行", "使用反馈、日志审计、提示词优化", "9A7BEA")]
    for i,(t,b,c) in enumerate(stages):
        x=.48+i*2.57; card(s,x,1.7,2.2,3.8,t,b,c,f"0{i+1}")
        if i<4: arrow(s,x+2.21,3.52,x+2.52,3.52,c)
    add_text(s,"建议验收指标：问题覆盖率｜接口调用正确率｜结果准确性｜响应时间｜权限合规｜用户使用频次",.75,6.15,11.85,.38,13,ORANGE,True,PP_ALIGN.CENTER)
    notes(s, "售前把演示问题沉淀为验收问题集，可直接连接产品、实施、测试和客户验收。")

    # 21 竞争优势
    s = new_slide(prs, 21, "我们的差异化，不只在模型", "模型可替换，真正形成壁垒的是行业数据、定位能力和交付方法")
    diffs=[("核心定位能力", "人与资产实时位置、轨迹、区域进出形成统一时空底座", ORANGE),
           ("行业工具体系", "面向医疗、安全、楼宇预置业务接口和回答规则", BLUE),
           ("数据隔离机制", "不同场景独立工具、数据和会话，减少串场和越权", TEAL),
           ("本地云端协同", "敏感查询本地化，复杂任务可按策略使用云端模型", GREEN),
           ("可演示可实施", "模拟接口快速验证，接口契约可平滑替换为真实系统", "9A7BEA"),
           ("销售到验收闭环", "用问题集驱动调研、试点、评测与验收", RED)]
    for i,(t,b,c) in enumerate(diffs):
        x=.62+(i%3)*4.18; y=1.45+(i//3)*2.35; card(s,x,y,3.76,1.78,t,b,c)
    notes(s, "竞争对比时不要只比模型名称和参数。把讨论拉回业务接口、数据治理、部署安全、定位基础和交付验收。")

    # 22 收口
    s = new_slide(prs, 22, "销售行动清单", "把一次产品演示转化为下一步可执行的项目动作")
    left=["确定客户主场景与关键决策人", "收集现有系统及接口清单", "确认 3—5 个演示黄金问题", "明确数据安全与部署边界"]
    right=["选定首期试点区域或科室", "形成可验收问题集与指标", "锁定数据、业务和信息化负责人", "约定接口调研与方案汇报时间"]
    card(s,.72,1.55,5.72,4.65,"演示前","\n".join("□ "+x for x in left),BLUE,"PREPARE")
    card(s,6.88,1.55,5.72,4.65,"演示后","\n".join("□ "+x for x in right),GREEN,"NEXT")
    add_text(s,"一句话收口：先用一个高价值场景验证，再复用统一平台扩展更多智能运营能力。",.78,6.45,11.75,.32,15,ORANGE,True,PP_ALIGN.CENTER)
    notes(s, "培训结尾要求每位销售选一个目标客户，完成场景、痛点、黄金问题、试点范围和下一步动作五项作业。")

    # 23-25 真实软件截图
    screenshots = [
        (23, "真实软件演示：医疗智能体", "切换后独立呈现医疗主题、能力入口和急诊/资产问题", "/mnt/c/Users/Administrator/Desktop/三智能体PPT-医疗界面.png", BLUE,
         "演示时先指出顶部智能体切换，再介绍下方医疗快捷问题和独立历史会话。"),
        (24, "真实软件演示：人员安全智能体", "围绕人员与资产定位、作业票、报警、考勤和能源设备展开", "/mnt/c/Users/Administrator/Desktop/三智能体PPT-人员安全界面.png", ORANGE,
         "这页重点展示人员安全场景最完整的能力入口，以及系统、Telegram、模型和工具调用状态。"),
        (25, "真实软件演示：政企楼宇智能体", "统一查询楼宇人员资产、能源空调、安防消防、停车电梯和维保", "/mnt/c/Users/Administrator/Desktop/三智能体PPT-楼宇界面.png", TEAL,
         "强调同一交互框架切换到楼宇后，主题、快捷问题、业务接口和会话均随场景改变。"),
    ]
    for no, title, subtitle, image_path, color, speaker_note in screenshots:
        s = new_slide(prs, no, title, subtitle)
        add_box(s, .7, 1.32, 11.93, 5.55, PANEL2, color)
        path = Path(image_path)
        if path.is_file():
            s.shapes.add_picture(str(path), Inches(.76), Inches(1.38), width=Inches(11.81), height=Inches(5.32))
        else:
            add_text(s, "截图文件未找到：" + str(path), 1.1, 3.6, 11.1, .4, 15, RED, True, PP_ALIGN.CENTER)
        notes(s, speaker_note)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)

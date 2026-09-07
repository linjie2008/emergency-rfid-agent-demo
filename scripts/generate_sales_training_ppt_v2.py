"""重构版：以清晰产品架构为主线生成工业产品销售培训 PPT。"""

from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from generate_sales_training_ppt import (
    BG, CYAN, FONT, GREEN, H, LINE, MUTED, OUTPUT, PANEL, PANEL_2, RED, STEEL,
    TEAL, W, WHITE, YELLOW, add_badge, add_box, add_capability_card, add_footer,
    add_metric, add_notes, add_rich_lines, add_text, add_title, new_slide, rgb,
    set_bg,
)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def pill(slide, text, x, y, w, color=TEAL):
    add_box(slide, x, y, w, 0.38, fill=PANEL_2, line=color)
    add_text(slide, text, x + 0.05, y + 0.08, w - 0.1, 0.19, size=9.5, color=color, bold=True, align=PP_ALIGN.CENTER)


def layer(slide, name, desc, x, y, w, color, number=""):
    add_box(slide, x, y, w, 0.72, fill=PANEL, line=color)
    if number:
        add_text(slide, number, x + 0.15, y + 0.19, 0.48, 0.24, size=10, color=color, bold=True)
    add_text(slide, name, x + (0.72 if number else 0.2), y + 0.17, 1.65, 0.28, size=13, color=color, bold=True)
    add_text(slide, desc, x + 2.05, y + 0.17, w - 2.3, 0.3, size=10.5, color=WHITE)


def app_card(slide, name, value, items, x, y, color):
    add_box(slide, x, y, 2.75, 2.25, fill=PANEL, line=color)
    add_text(slide, name, x + 0.2, y + 0.22, 2.35, 0.32, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, value, x + 0.22, y + 0.72, 2.3, 0.46, size=10.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, items, x + 0.26, y + 1.32, 2.22, 0.7, size=9.5, color=MUTED, bullet=True, spacing=5)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "人员安全智能运营中心｜工业产品销售培训"
    prs.core_properties.subject = "1个中心、2个底座、4大应用、N个入口"
    prs.core_properties.author = "人员安全智能运营中心"

    # 01 封面
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    add_box(slide, 0, 0, 0.16, H, fill=YELLOW, line=YELLOW, radius=False, width=0)
    for i in range(8):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.2 + i), Inches(0), Inches(2.6 + i), Inches(7.5))
        line.line.color.rgb = rgb("142834"); line.line.width = Pt(0.7)
    add_text(slide, "人员安全智能运营中心", 0.85, 1.65, 9.5, 0.76, size=36, bold=True)
    add_text(slide, "面向工业现场的安全与智能运营平台", 0.88, 2.65, 8.7, 0.52, size=22, color=YELLOW, bold=True)
    add_text(slide, "连接人员、空间、作业、设备与生产数据\n形成感知—分析—决策—处置—复盘的业务闭环", 0.9, 3.6, 7.8, 1.05, size=17, color="C0D0D7")
    add_badge(slide, "工业产品销售培训", 0.9, 5.3, 1.9, YELLOW)
    add_badge(slide, "产品架构版", 3.0, 5.3, 1.45, TEAL)
    add_text(slide, "INDUSTRIAL SAFETY INTELLIGENCE", 0.92, 6.65, 4.6, 0.25, size=9, color="647B86")
    add_notes(slide, "开场定位：这不是一个聊天机器人，也不是新的数据孤岛，而是连接工业现场既有系统的安全智能运营中心。")

    # 02 逻辑起点
    slide = new_slide(prs, 2, "产品逻辑起点：客户缺的不是系统，而是统一运营能力", "多数工业企业已经完成系统建设，但数据仍停留在各自系统内部")
    left = [("人员定位", "知道坐标，不知道是否合规"), ("作业票", "知道许可，不知道现场实际行为"), ("设备监测", "知道当前值，不知道未来风险"), ("生产系统", "知道运行参数，缺少跨专业关联")]
    for i, (title, desc) in enumerate(left):
        y = 1.5 + i * 1.12
        add_box(slide, 0.72, y, 4.45, 0.82, fill=PANEL, line=LINE)
        add_text(slide, title, 0.96, y + 0.2, 1.25, 0.27, size=13, color=TEAL, bold=True)
        add_text(slide, desc, 2.25, y + 0.2, 2.6, 0.35, size=11, color=MUTED)
    arrow(slide, 5.35, 3.6, 6.35, 3.6, YELLOW, 2.5)
    add_box(slide, 6.55, 1.5, 5.95, 4.18, fill=PANEL_2, line=YELLOW)
    add_text(slide, "人员安全智能运营中心", 7.0, 1.9, 5.05, 0.42, size=22, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "统一对象", 7.18, 2.75, 1.25, 0.3, size=13, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "统一入口", 8.66, 2.75, 1.25, 0.3, size=13, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "统一闭环", 10.15, 2.75, 1.25, 0.3, size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "把分散数据转化为可查询、可关联、可预测、可执行的现场运营能力", 7.1, 3.55, 4.85, 1.05, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "销售主张：不替换现有系统，释放现有系统的数据价值", 0.75, 6.32, 11.75, 0.35, size=15, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "从客户已有系统讲起。定位系统、票务系统和设备系统都有价值，但彼此独立；产品的角色是连接、理解和闭环。")

    # 03 产品公式
    slide = new_slide(prs, 3, "一张图讲清产品：1 + 2 + 4 + N", "所有销售材料、方案和演示都围绕这一结构展开")
    blocks = [
        ("1", "一个运营中心", "统一态势、统一查询、统一决策", YELLOW, 0.7, 3.0),
        ("2", "两个技术底座", "工业数据底座 + AI 智能体底座", TEAL, 3.95, 3.0),
        ("4", "四大业务应用", "人员安全｜作业安全｜设备健康｜能源运营", CYAN, 7.2, 3.0),
        ("N", "多种使用入口", "PC｜大屏｜移动端｜开放 API", GREEN, 10.45, 2.15),
    ]
    for num, title, desc, color, x, w in blocks:
        add_box(slide, x, 1.55, w, 3.8, fill=PANEL, line=color)
        add_text(slide, num, x + 0.2, 1.9, w - 0.4, 0.88, size=44, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.2, 3.05, w - 0.4, 0.38, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.25, 3.78, w - 0.5, 0.88, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.75, 5.85, 11.85, 0.62, fill=PANEL_2, line=LINE)
    add_text(slide, "统一产品内核，按客户场景组合业务应用；从一个场景切入，逐步扩展", 1.0, 6.04, 11.3, 0.25, size=13.5, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "这是整套 PPT 的导航页。销售必须能脱稿讲清 1+2+4+N，后面每一部分都是对这张图的展开。")

    # 04 核心架构
    slide = new_slide(prs, 4, "产品总体架构：从现场数据到业务闭环", "纵向五层架构，横向安全治理与运行评测")
    # 左侧层级
    layer(slide, "使用入口", "PC 工作台｜指挥大屏｜Telegram 移动端｜第三方系统 API", 0.65, 1.38, 10.55, YELLOW, "L5")
    layer(slide, "业务应用", "人员空间安全｜作业合规管理｜设备预测性维护｜能源运营分析", 0.65, 2.22, 10.55, CYAN, "L4")
    layer(slide, "AI 智能体底座", "意图识别｜动态工具路由｜知识检索｜上下文管理｜结果生成", 0.65, 3.06, 10.55, TEAL, "L3")
    layer(slide, "工业数据底座", "统一对象模型｜主数据映射｜实时/历史数据服务｜指标口径", 0.65, 3.9, 10.55, GREEN, "L2")
    layer(slide, "连接与现场", "定位/门禁/作业票｜DCS/SCADA｜EAM/CMMS｜IoT 终端｜数据库", 0.65, 4.74, 10.55, "7B8E98", "L1")
    # 右侧治理条
    add_box(slide, 11.45, 1.38, 1.2, 4.08, fill="211B16", line=YELLOW)
    add_text(slide, "安\n全\n治\n理\n与\n运\n行\n评\n测", 11.82, 1.64, 0.45, 3.45, size=12, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.65, 5.82, 12.0, 0.58, fill=PANEL_2, line=LINE)
    add_text(slide, "权限｜脱敏｜审计｜模型评测｜工具调用追踪｜延迟与成本监控｜告警闭环", 0.92, 6.0, 11.45, 0.24, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_notes(slide, "架构讲法从下往上：接入现有系统，形成统一数据对象，智能体理解问题并调用接口，四大应用形成业务价值，最终通过多入口服务不同角色。安全治理贯穿全栈。")

    # 05 两个底座
    slide = new_slide(prs, 5, "两个底座决定产品是否可持续", "业务应用可以组合，但数据语义与智能编排必须统一")
    add_box(slide, 0.72, 1.5, 5.75, 4.8, fill=PANEL, line=GREEN)
    add_text(slide, "工业数据底座", 1.05, 1.85, 5.05, 0.42, size=23, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "解决“数据如何统一”", 1.1, 2.45, 4.95, 0.3, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, ["统一人、空间、作业、设备、事件对象", "适配 API、数据库、消息和时序数据", "统一编码、时间、组织和指标口径", "为业务关联和权限控制提供依据"], 1.12, 3.1, 4.95, 2.25, size=13, color=MUTED, bullet=True, spacing=14)
    add_box(slide, 6.85, 1.5, 5.75, 4.8, fill=PANEL, line=TEAL)
    add_text(slide, "AI 智能体底座", 7.18, 1.85, 5.05, 0.42, size=23, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "解决“问题如何变成行动”", 7.2, 2.45, 5.0, 0.3, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, ["识别业务意图并选择相关接口组", "按权限执行查询、分析和可视化工具", "结合制度知识和历史上下文生成结论", "记录工具、延迟、tokens、成本与质量"], 7.25, 3.1, 4.9, 2.25, size=13, color=MUTED, bullet=True, spacing=14)
    add_notes(slide, "两个底座分别回答数据问题和智能问题。客户如果只买某个应用，也复用同一底座，为后续扩展降低成本。")

    # 06 统一对象模型
    slide = new_slide(prs, 6, "统一业务对象：产品不是接口集合", "通过五类核心对象建立跨系统关联关系")
    objects = [("人", "员工/承包商/访客", YELLOW), ("空间", "厂区/建筑/区域", TEAL), ("作业", "作业票/工序/工时", CYAN), ("设备", "资产/终端/部件", GREEN), ("事件", "报警/缺陷/处置", RED)]
    for i, (name, desc, color) in enumerate(objects):
        x = 0.55 + i * 2.52
        add_box(slide, x, 1.52, 2.12, 1.25, fill=PANEL, line=color)
        add_text(slide, name, x + 0.15, 1.76, 1.82, 0.35, size=19, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.15, 2.25, 1.82, 0.22, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.75, 3.32, 11.8, 2.4, fill=PANEL_2, line=LINE)
    relations = ["谁", "在什么位置", "执行什么作业", "使用/维护什么设备", "发生什么事件"]
    colors = [YELLOW, TEAL, CYAN, GREEN, RED]
    for i, (text, color) in enumerate(zip(relations, colors)):
        x = 1.0 + i * 2.25
        pill(slide, text, x, 4.0, 1.75, color)
        if i < 4: add_text(slide, "→", x + 1.82, 4.05, 0.35, 0.2, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "关联后才能回答：某人员是否在许可时间进入许可区域，并对目标设备执行合规作业？", 1.0, 4.85, 10.95, 0.38, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "接口是技术实现，对象关系才是产品资产。这一页用于解释为什么平台能跨系统回答问题。")

    # 07 运营闭环
    slide = new_slide(prs, 7, "统一运营闭环：从“看见”到“处置”", "所有应用共享同一条闭环，不再停留在查询和报表")
    stages = [("感知", "采集现场事实", CYAN), ("连接", "形成统一对象", GREEN), ("分析", "发现异常与关联", TEAL), ("决策", "给出优先级与建议", YELLOW), ("执行", "触发处置与工单", RED), ("复盘", "沉淀规则与指标", "9B84E8")]
    for i, (name, desc, color) in enumerate(stages):
        x = 0.45 + i * 2.12
        add_box(slide, x, 1.65, 1.72, 1.55, fill=PANEL, line=color)
        add_text(slide, name, x + 0.15, 1.92, 1.42, 0.33, size=17, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.16, 2.48, 1.4, 0.35, size=10, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 5: add_text(slide, "→", x + 1.76, 2.22, 0.3, 0.25, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.72, 3.85, 11.9, 1.55, fill=PANEL_2, line=LINE)
    add_text(slide, "产品输出", 1.0, 4.15, 1.25, 0.3, size=15, color=YELLOW, bold=True)
    add_text(slide, "实时态势｜异常清单｜原因证据｜风险等级｜处置建议｜闭环状态｜复盘指标", 2.35, 4.13, 9.7, 0.38, size=13.5, color=WHITE, bold=True)
    add_text(slide, "查询只是入口，闭环才是最终价值", 0.75, 6.15, 11.85, 0.36, size=16, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "销售讲每个功能时都要回到闭环：发现后谁处置、如何跟踪、如何复盘。只有查询没有闭环，很难形成持续预算。")

    # 08 四大应用总览
    slide = new_slide(prs, 8, "四大业务应用，共享两个底座", "根据客户痛点组合交付，不要求一次覆盖全部范围")
    app_card(slide, "人员空间安全", "回答“谁在现场、现在在哪、是否异常”", ["人员与组织", "定位与轨迹", "门禁/考勤/驻留"], 0.65, 1.55, YELLOW)
    app_card(slide, "作业安全合规", "回答“谁按什么许可在什么位置作业”", ["作业票合规", "流程与工时", "轨迹关联"], 3.72, 1.55, CYAN)
    app_card(slide, "设备预测维护", "回答“哪台设备可能出问题、何时处理”", ["健康度", "故障预测/RUL", "维护建议"], 6.79, 1.55, GREEN)
    app_card(slide, "能源运营分析", "回答“电热气运行如何、风险在哪里”", ["生产与负荷", "调度与能效", "能源安全"], 9.86, 1.55, TEAL)
    add_box(slide, 0.72, 4.55, 11.85, 1.3, fill=PANEL_2, line=LINE)
    add_text(slide, "公共能力", 0.98, 4.86, 1.15, 0.3, size=14, color=YELLOW, bold=True)
    common = ["透明建筑数字孪生", "表格与图表", "知识库", "多模型", "权限审计", "运行评测"]
    for i, name in enumerate(common): pill(slide, name, 2.25 + i * 1.62, 4.78, 1.38, TEAL if i < 3 else CYAN)
    add_notes(slide, "四大应用是销售和报价的模块边界；公共能力由平台统一提供。客户通常从人员空间安全或设备预测维护切入。")

    # 09 人员应用
    slide = new_slide(prs, 9, "应用一｜人员空间安全", "以人员身份为主线，连接准入、位置、行为与事件")
    chain = [("身份", "员工/承包商/访客"), ("准入", "资质/培训/门禁"), ("位置", "实时/轨迹/进出"), ("行为", "驻留/距离/强度"), ("风险", "报警/处置/复盘")]
    for i, (name, desc) in enumerate(chain):
        x = 0.55 + i * 2.5
        add_box(slide, x, 1.55, 2.08, 1.25, fill=PANEL, line=TEAL if i < 4 else RED)
        add_text(slide, name, x + 0.15, 1.82, 1.78, 0.3, size=16, color=TEAL if i < 4 else RED, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.15, 2.3, 1.78, 0.24, size=9.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.72, 3.45, 12.0, 2.25, fill=PANEL_2, line=LINE)
    add_text(slide, "典型问题", 1.02, 3.8, 1.25, 0.3, size=14, color=YELLOW, bold=True)
    add_rich_lines(slide, ["当前受限空间作业区有哪些承包商？", "张伟今天经过哪些区域，停留多久？", "哪些人员存在连续报警或疲劳风险？"], 2.4, 3.75, 5.1, 1.45, size=13, color=WHITE, bullet=True, spacing=11)
    add_text(slide, "业务价值", 8.1, 3.8, 1.25, 0.3, size=14, color=TEAL, bold=True)
    add_rich_lines(slide, ["缩短人员查找时间", "强化承包商全过程管理", "提升应急人员清点效率"], 9.45, 3.75, 2.8, 1.45, size=12, color=WHITE, bullet=True, spacing=11)
    add_notes(slide, "人员定位只是应用中的一环。销售重点讲身份—准入—位置—行为—风险的完整链路。")

    # 10 作业应用
    slide = new_slide(prs, 10, "应用二｜作业安全合规", "把票面许可与人员实际行为关联，形成可追溯证据链")
    add_box(slide, 0.72, 1.5, 3.15, 4.7, fill=PANEL, line=CYAN)
    add_text(slide, "票面要求", 1.05, 1.88, 2.5, 0.38, size=20, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, ["作业类型", "许可时间", "许可区域", "责任人与人员", "安全措施"], 1.18, 2.65, 2.2, 2.4, size=13, color=WHITE, bullet=True, spacing=15)
    arrow(slide, 4.08, 3.82, 5.05, 3.82, YELLOW, 2)
    add_box(slide, 5.18, 1.5, 3.15, 4.7, fill=PANEL, line=YELLOW)
    add_text(slide, "现场事实", 5.5, 1.88, 2.5, 0.38, size=20, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, ["人员到场", "区域进出", "停留时长", "报警事件", "有效工时"], 5.65, 2.65, 2.2, 2.4, size=13, color=WHITE, bullet=True, spacing=15)
    arrow(slide, 8.55, 3.82, 9.45, 3.82, TEAL, 2)
    add_box(slide, 9.58, 1.5, 2.95, 4.7, fill=PANEL_2, line=TEAL)
    add_text(slide, "合规输出", 9.88, 1.88, 2.35, 0.38, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_rich_lines(slide, ["单票报告", "合规趋势", "违规排行", "流程时长", "整改闭环"], 10.0, 2.65, 2.0, 2.4, size=13, color=WHITE, bullet=True, spacing=15)
    add_notes(slide, "这页讲清作业安全应用的输入、关联和输出。客户已有作业票系统时，我们补的是现场事实和综合分析。")

    # 11 设备应用
    slide = new_slide(prs, 11, "应用三｜设备预测性维护", "监测终端只是感知入口，最终目标是优化维护决策")
    stages = [("采集", "振动/温度/电流\n局放/油色谱", CYAN), ("健康评价", "健康度\n异常分数", GREEN), ("风险预测", "故障模式\n发生概率", YELLOW), ("维护决策", "RUL/窗口\n备件/工单", RED)]
    for i, (name, desc, color) in enumerate(stages):
        x = 0.72 + i * 3.05
        add_box(slide, x, 1.55, 2.65, 1.72, fill=PANEL, line=color)
        add_text(slide, name, x + 0.18, 1.85, 2.3, 0.32, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.2, 2.38, 2.25, 0.52, size=11, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3: add_text(slide, "→", x + 2.7, 2.16, 0.3, 0.25, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    add_box(slide, 0.72, 3.85, 12.0, 1.72, fill=PANEL_2, line=LINE)
    add_text(slide, "输出必须可解释", 1.02, 4.18, 1.8, 0.32, size=15, color=YELLOW, bold=True)
    add_text(slide, "设备｜健康评分｜异常趋势｜故障模式｜预测概率｜模型置信度｜建议窗口｜建议动作", 3.0, 4.15, 9.1, 0.4, size=12.5, color=WHITE, bold=True)
    add_text(slide, "产品边界：辅助运维决策，不替代保护系统、联锁系统和专业诊断", 1.02, 5.02, 10.95, 0.27, size=11, color=RED, bold=True)
    add_notes(slide, "预测维护的产品结构是感知—评价—预测—决策，不要把产品讲成传感器平台。销售必须说明可解释性和专业复核边界。")

    # 12 能源应用
    slide = new_slide(prs, 12, "应用四｜能源运营分析", "面向综合能源企业的跨专业运行与风险视角")
    domains = [("电", "机组出力｜电网负荷｜变电站｜线路｜停电", YELLOW), ("热", "供回温｜压力｜流量｜热负荷｜热网报警", RED), ("气", "门站｜调压站｜供气量｜压力｜加臭状态", TEAL)]
    for i, (name, desc, color) in enumerate(domains):
        x = 0.72 + i * 4.08
        add_box(slide, x, 1.55, 3.62, 2.15, fill=PANEL, line=color)
        add_text(slide, name, x + 0.25, 1.9, 0.72, 0.62, size=32, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 1.1, 1.93, 2.15, 1.05, size=11.5, color=WHITE)
    add_box(slide, 0.72, 4.25, 11.78, 1.4, fill=PANEL_2, line=LINE)
    add_text(slide, "统一运营视角", 1.0, 4.58, 1.6, 0.32, size=15, color=YELLOW, bold=True)
    add_text(slide, "生产状态 + 调度执行 + 能效指标 + 设备风险 + 人员安全", 2.9, 4.55, 8.7, 0.38, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "能源运营应用适用于电热气综合客户。单一工业客户可替换为其生产专业数据，架构不变。")

    # 13 公共空间能力
    slide = new_slide(prs, 13, "公共能力｜透明建筑数字孪生", "数字孪生是空间呈现能力，为人员、作业、设备和事件提供统一位置视角")
    add_box(slide, 0.65, 1.45, 7.4, 4.95, fill="111A20", line=LINE)
    # 沙盘
    for y in (3.02, 5.05):
        add_box(slide, 0.92, y, 6.85, 0.25, fill="2B353A", line="2B353A", radius=False, width=0)
    buildings = [(1.05,1.75,1.7,1.02,"中控楼"),(3.05,1.65,2.1,1.15,"装置区"),(5.52,1.78,1.72,1.02,"配电室"),(1.18,3.55,2.05,1.05,"检修区"),(3.72,3.45,1.45,1.18,"罐区"),(5.62,3.6,1.55,1.02,"泵房")]
    for i,(x,y,w,h,name) in enumerate(buildings):
        add_box(slide,x,y,w,h,fill="233841",line=TEAL if i in (1,5) else "607782",radius=False,width=1.2)
        add_text(slide,name,x+0.08,y+h-0.24,w-0.16,0.17,size=8,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    for x,y,name in [(3.62,2.1,"张伟"),(4.25,2.4,"李强"),(6.18,4.05,"王磊")]:
        dot=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.18),Inches(.18)); dot.fill.solid(); dot.fill.fore_color.rgb=rgb(GREEN); dot.line.color.rgb=rgb(WHITE)
        add_text(slide,name,x-.22,y-.25,.62,.18,size=7,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    add_text(slide, "人在透明建筑内部显示", 1.1, 5.65, 6.45, 0.3, size=12, color=YELLOW, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "统一空间视角", 8.55, 1.6, 2.8, 0.38, size=19, color=YELLOW, bold=True)
    add_rich_lines(slide, ["人员实时位置", "作业许可区域", "设备与监测终端", "报警与风险区域", "道路、建筑与空间标签"], 8.55, 2.3, 3.55, 2.8, size=14, color=WHITE, bullet=True, spacing=15)
    pill(slide, "可对接 RFID / UWB / AoA / GPS", 8.55, 5.5, 3.45, TEAL)
    add_notes(slide, "明确数字孪生在架构中的位置：它是公共空间呈现能力，不是孤立模块。定位精度由现场定位系统决定。")

    # 14 智能体架构
    slide = new_slide(prs, 14, "AI 智能体工作架构", "动态路由、接口执行和结果治理构成可信问答链路")
    flow = [("问题理解", "识别对象/时间/意图", YELLOW), ("业务路由", "选择相关应用域", TEAL), ("工具编排", "加载少量相关接口", CYAN), ("数据执行", "按权限查询真实数据", GREEN), ("结果治理", "校验/脱敏/格式化", RED), ("多端输出", "文本/表格/图表/地图", "9B84E8")]
    for i,(name,desc,color) in enumerate(flow):
        x=.42+i*2.12
        add_box(slide,x,1.55,1.76,1.55,fill=PANEL,line=color)
        add_text(slide,name,x+.1,1.82,1.56,.3,size=14,color=color,bold=True,align=PP_ALIGN.CENTER)
        add_text(slide,desc,x+.12,2.35,1.52,.42,size=9.5,color=MUTED,align=PP_ALIGN.CENTER)
        if i<5:add_text(slide,"→",x+1.79,2.15,.27,.22,size=14,color=MUTED,align=PP_ALIGN.CENTER)
    add_box(slide,.72,3.75,12,1.82,fill=PANEL_2,line=LINE)
    add_text(slide,"动态工具路由",1.0,4.08,1.8,.32,size=15,color=YELLOW,bold=True)
    add_text(slide,"60 个业务接口全部可访问，但每次只向模型加载相关接口组",3.0,4.05,8.7,.36,size=15,color=WHITE,bold=True)
    add_text(slide,"效果：降低上下文占用｜减少工具误选｜提高本地模型速度｜支持跨域组合与连续追问",1.0,4.82,10.9,.3,size=11.5,color=TEAL)
    add_notes(slide, "这页回答产品如何避免模型幻觉和工具过多问题。强调先路由、再执行、后治理；业务数据来自接口，不由模型编造。")

    # 15 部署治理
    slide = new_slide(prs, 15, "部署与治理架构", "本地优先、云端可选；根据数据分级决定模型和网络边界")
    add_box(slide,.7,1.48,5.8,4.55,fill=PANEL,line=TEAL)
    add_text(slide,"企业内网区",1.0,1.82,5.2,.38,size=21,color=TEAL,bold=True,align=PP_ALIGN.CENTER)
    add_rich_lines(slide,["业务接口与统一数据服务","本地 Qwen 模型与 GPU 推理","智能体服务、知识库和权限","人员与生产敏感数据留在内网"],1.1,2.55,4.9,2.3,size=13,color=WHITE,bullet=True,spacing=15)
    add_box(slide,6.85,1.48,5.8,4.55,fill=PANEL,line=CYAN)
    add_text(slide,"可选云端能力",7.15,1.82,5.2,.38,size=21,color=CYAN,bold=True,align=PP_ALIGN.CENTER)
    add_rich_lines(slide,["复杂推理与长报告","多模型灵活切换","只发送经过策略允许的数据","网络异常自动重试与节点回退"],7.25,2.55,4.9,2.3,size=13,color=WHITE,bullet=True,spacing=15)
    add_box(slide,.72,6.25,11.9,.42,fill="241D12",line=YELLOW)
    add_text(slide,"横向治理：身份权限｜字段脱敏｜操作审计｜模型评测｜tokens/成本｜运行监控",1.0,6.35,11.3,.2,size=10.5,color=YELLOW,bold=True,align=PP_ALIGN.CENTER)
    add_notes(slide,"不要把本地和云端讲成二选一。按数据分级和任务复杂度组合；客户要求完全内网时可只启用本地模型。")

    # 16 三大组合场景
    slide = new_slide(prs, 16, "三类高价值组合场景", "组合场景比单接口查询更能体现产品价值")
    scenarios=[("高风险作业闭环","人员 + 空间 + 作业票 + 报警","确认谁在许可时间进入许可区域，是否发生异常并完成处置",YELLOW),("设备风险提前识别","监测终端 + 设备健康 + 工单","识别异常趋势、预测故障并推荐维护窗口",GREEN),("综合运营态势","人员安全 + 设备风险 + 能源运行","管理层在一个入口掌握安全、生产和设备重点",TEAL)]
    for i,(name,combo,outcome,color) in enumerate(scenarios):
        x=.65+i*4.18
        add_box(slide,x,1.55,3.75,4.55,fill=PANEL,line=color)
        add_text(slide,f"0{i+1}",x+.2,1.85,.5,.3,size=11,color=color,bold=True)
        add_text(slide,name,x+.25,2.35,3.25,.52,size=19,color=color,bold=True,align=PP_ALIGN.CENTER)
        pill(slide,combo,x+.35,3.18,3.05,color)
        add_text(slide,outcome,x+.38,4.0,3.0,1.2,size=12.5,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    add_notes(slide, "销售优先卖组合场景。单点功能容易与客户现有系统比较，组合场景体现跨系统关联和智能运营价值。")

    # 17 产品组合
    slide = new_slide(prs, 17, "产品组合与商业边界", "平台底座统一，业务应用按客户范围组合")
    add_box(slide,.72,1.45,12,1.02,fill=PANEL_2,line=YELLOW)
    add_text(slide,"基础平台",1.0,1.78,1.45,.32,size=17,color=YELLOW,bold=True)
    add_text(slide,"统一门户｜AI 智能体｜数据连接｜权限审计｜报表图表｜运行监控",2.65,1.76,9.35,.38,size=13,color=WHITE)
    modules=[("人员空间安全",YELLOW),("作业安全合规",CYAN),("设备预测维护",GREEN),("能源运营分析",TEAL)]
    for i,(name,color) in enumerate(modules):
        x=.72+i*3.05
        add_box(slide,x,2.9,2.7,1.55,fill=PANEL,line=color)
        add_text(slide,name,x+.2,3.3,2.3,.35,size=15,color=color,bold=True,align=PP_ALIGN.CENTER)
        add_text(slide,"业务应用模块",x+.2,3.85,2.3,.25,size=9.5,color=MUTED,align=PP_ALIGN.CENTER)
    add_box(slide,.72,4.9,12,1.0,fill=PANEL_2,line=LINE)
    add_text(slide,"项目服务",1.0,5.23,1.45,.32,size=16,color=TEAL,bold=True)
    add_text(slide,"数据适配｜现场建模｜算法验证｜系统集成｜培训运维｜持续评测",2.65,5.22,9.35,.35,size=12.5,color=WHITE)
    add_text(slide,"报价逻辑：基础平台 + 应用模块 + 数据/集成服务 + 可选模型与硬件资源",.75,6.35,11.85,.3,size=13,color=YELLOW,bold=True,align=PP_ALIGN.CENTER)
    add_notes(slide, "这页用于帮助销售建立报价结构。不要按接口数量报价，应按平台、应用、数据适配、集成范围和部署资源拆分。")

    # 18 客户价值与KPI
    slide = new_slide(prs, 18, "客户价值与验收指标", "先建立现状基线，再与客户共同确定目标值")
    values=[("效率","跨系统查询时间\n临时报表时间",YELLOW),("安全","报警闭环时长\n异常发现覆盖率",RED),("合规","作业票合规率\n高频违规下降",CYAN),("设备","异常提前量\n非计划停机时长",GREEN),("运营","数据使用频次\n管理决策响应",TEAL)]
    for i,(name,kpi,color) in enumerate(values):
        x=.48+i*2.53
        add_box(slide,x,1.55,2.15,2.55,fill=PANEL,line=color)
        add_text(slide,name,x+.18,1.9,1.8,.42,size=22,color=color,bold=True,align=PP_ALIGN.CENTER)
        add_text(slide,kpi,x+.2,2.75,1.75,.8,size=11.5,color=WHITE,align=PP_ALIGN.CENTER)
    add_box(slide,.72,4.65,12,1.15,fill=PANEL_2,line=LINE)
    add_text(slide,"销售原则",1.02,4.98,1.35,.32,size=15,color=YELLOW,bold=True)
    add_text(slide,"不在缺少基线数据时承诺固定百分比；以真实接口、真实场景和共同定义的验收口径验证价值",2.55,4.9,9.45,.52,size=12.5,color=WHITE,bold=True)
    add_notes(slide, "价值必须可量化，但目标值由客户基线决定。推荐在试点前记录现状流程耗时和风险指标。")

    # 19 销售打法
    slide = new_slide(prs, 19, "销售打法：从一个高价值场景进入", "场景明确、数据可得、责任部门清晰、价值可以量化")
    steps=[("找痛点","哪项管理工作最慢、风险最高？",YELLOW),("定场景","选择人员作业或重点设备",TEAL),("查数据","确认系统、接口、历史数据",CYAN),("定指标","确定现状基线与验收口径",GREEN),("做闭环","从查询延伸到处置与复盘",RED)]
    for i,(name,desc,color) in enumerate(steps):
        x=.52+i*2.53
        add_box(slide,x,1.55,2.13,1.65,fill=PANEL,line=color)
        add_text(slide,name,x+.15,1.85,1.83,.32,size=16,color=color,bold=True,align=PP_ALIGN.CENTER)
        add_text(slide,desc,x+.18,2.4,1.78,.44,size=10,color=MUTED,align=PP_ALIGN.CENTER)
    add_box(slide,.72,3.85,12,1.85,fill=PANEL_2,line=LINE)
    add_text(slide,"优先切入场景",1.02,4.18,1.65,.32,size=15,color=YELLOW,bold=True)
    add_rich_lines(slide,["高风险区域人员与作业票联动","承包商全过程安全管理","重点设备异常趋势与维护窗口"],2.95,4.1,4.15,1.25,size=12.5,color=WHITE,bullet=True,spacing=10)
    add_text(slide,"关键参与人",8.0,4.18,1.35,.32,size=15,color=TEAL,bold=True)
    add_rich_lines(slide,["业务责任部门", "设备/安全专业人员", "信息化与数据团队"],9.42,4.1,2.8,1.25,size=12,color=WHITE,bullet=True,spacing=10)
    add_notes(slide, "不要一开始推全平台。选一个痛点强、数据基础好的闭环场景，验证后用统一底座扩展。")

    # 20 标准演示与架构回收
    slide = new_slide(prs, 20, "标准销售演示：始终回到产品架构", "15 分钟内让客户理解平台、应用、数据和价值")
    demo=[("2分钟","产品公式","1个中心 + 2个底座 + 4大应用 + N个入口"),("3分钟","人员空间","自然语言查询 → 定位接口 → 透明三维建筑"),("3分钟","作业合规","作业票 → 人员轨迹 → 合规结论与报表"),("3分钟","设备维护","监测终端 → 风险预测 → 维护建议"),("2分钟","技术底座","动态工具路由 → 本地模型 → token统计"),("2分钟","下一步","确认一个场景、接口清单和价值指标")]
    for i,(time,name,desc) in enumerate(demo):
        y=1.38+i*.85
        pill(slide,time,.78,y+.1,.9,TEAL)
        add_text(slide,name,1.95,y+.14,1.35,.28,size=12.5,color=WHITE,bold=True)
        add_box(slide,3.45,y,8.85,.58,fill=PANEL,line=LINE)
        add_text(slide,desc,3.72,y+.16,8.3,.24,size=10.8,color=MUTED)
    add_notes(slide, "演示不从模型配置开始。先讲产品公式，再演示两个业务闭环，最后说明底座和下一步。")

    # 21 异议
    slide = new_slide(prs, 21, "常见异议：用架构回答，而不是只讲模型", "承认边界、解释机制、给出验证方式")
    objections=[("已有定位/作业票系统","保留原系统；通过数据底座统一对象，通过业务应用形成跨系统闭环。"),("大模型会编造","业务数据必须来自工具接口；结果经过权限、脱敏和格式治理，并记录调用过程。"),("敏感数据不能上云","支持本地模型和内网部署；云端能力按数据分级策略选择。"),("本地模型能力有限","动态路由每次只加载相关接口；复杂任务可分段本地处理或按策略使用云端。"),("预测结果不可靠","展示数据证据、趋势和置信度，通过历史回测和现场试点确定使用边界。")]
    for i,(q,a) in enumerate(objections):
        y=1.38+i*1.05
        add_box(slide,.72,y,3.65,.8,fill="21191A",line=RED)
        add_text(slide,q,.95,y+.2,3.2,.32,size=11,color="FFB4B4",bold=True)
        add_box(slide,4.6,y,7.95,.8,fill=PANEL,line=TEAL)
        add_text(slide,a,4.86,y+.15,7.43,.48,size=10.5,color=WHITE)
    add_notes(slide, "异议处理围绕整体架构：原系统在连接层，统一对象在数据底座，模型在智能体底座，价值在业务应用。")

    # 22 收尾
    slide=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide)
    add_box(slide,0,0,.16,H,fill=YELLOW,line=YELLOW,radius=False,width=0)
    add_text(slide,"人员安全智能运营中心",.85,1.2,8.5,.62,size=32,bold=True)
    add_text(slide,"1 个中心 · 2 个底座 · 4 大应用 · N 个入口",.88,2.15,9.5,.52,size=23,color=YELLOW,bold=True)
    add_text(slide,"以统一数据连接现场\n以智能体驱动决策\n以业务闭环创造价值",.9,3.25,5.7,1.45,size=20,color="C4D2D8",bold=True)
    add_box(slide,8.6,1.42,3.35,3.35,fill=PANEL,line=TEAL)
    add_text(slide,"下一步",9.25,1.9,2.05,.4,size=20,color=TEAL,bold=True,align=PP_ALIGN.CENTER)
    add_text(slide,"选择一个场景\n盘点一组接口\n定义一项指标\n完成一次价值验证",9.15,2.72,2.25,1.45,size=15,color=WHITE,bold=True,align=PP_ALIGN.CENTER)
    add_badge(slide,"工业安全 · 设备健康 · 智能运营",.9,5.65,3.6,YELLOW)
    add_notes(slide, "以行动收尾：邀请客户安排业务、设备和信息化人员，围绕一个高价值场景开展数据诊断。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"generated={OUTPUT}")
    print(f"slides={len(prs.slides)}")


if __name__ == "__main__":
    build()
